# -*- coding: utf-8 -*-
"""產出後的硬規則檢查（gen_pptx.py 會自動呼叫）。

用法：python check_deck.py --course immune --pptx <檔案> [--deck dent] [--section N]
檢查項（門檻都在 course.yaml 的 sizes）：
  1. 張數＝選單展開後張數；畫幅＝slide.width×height
  2. 所有 run 的 latin／ea 字型＝Microsoft JhengHei
  3. 字級：出處（「圖：」「資料來源：」開頭）≥ credit，其餘 ≥ min（RED）；內文 < body_warn 列 WARN
  4. 每張可見字元：> char_red RED、> char_warn WARN（封面、divider、國考題不算）
  5. 有圖率：內容張（排除封面、divider、國考題）有真圖或框架圖的比例 ≥ img_rate_min（RED）
  6. 殘留佔位字（圖片待補、▦）與靜態頁碼：RED
  7. 國考題 key 必須在該 deck 的題庫；各節張數在 per_section 範圍（WARN）
  8. gen_pptx 的 _layout.json：估算溢出 >15% RED、其餘溢出 WARN；渲染錯誤 RED
"""
import argparse
import json
import os
import re
import sys

from pptx import Presentation
from pptx.util import Emu

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import FONT, build_deck, load_course, visible_text  # noqa: E402

A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
CREDIT_PREFIX = ("圖：", "圖:", "資料來源", "圖片來源", "圖源")
NO_COUNT = {"title", "divider", "quiz"}


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


def iter_paras(slide):
    for sh in walk(slide.shapes):
        frames = []
        if sh.has_text_frame:
            frames.append(sh.text_frame)
        if getattr(sh, "has_table", False) and sh.has_table:
            frames += [c.text_frame for row in sh.table.rows for c in row.cells]
        for tf in frames:
            for p in tf.paragraphs:
                yield p


def main(course, pptx_path, deck, section=None):
    S = course.sizes
    ok = True
    reds, warns = [], []
    built = build_deck(course, deck)
    if section:  # 與 gen_pptx.build_pptx 相同：只留該節並從 1 重新編號，張號才對得上 PPTX
        lo, hi, t = built["sections"][section - 1]
        built["slides"] = [s for s in built["slides"] if lo <= s["num"] <= hi]
        for k, s in enumerate(built["slides"], 1):
            s["num"] = k
        built["sections"] = [(1, len(built["slides"]), t)]
    prs = Presentation(pptx_path)
    n = len(prs.slides)
    print(f"\n===== check_deck [{built['name']}] {os.path.basename(pptx_path)} =====")
    if n != len(built["slides"]):
        reds.append(f"張數 {n} ≠ 選單展開 {len(built['slides'])}")
    else:
        print(f"[ok] 張數 {n}")
    w, h = Emu(prs.slide_width).inches, Emu(prs.slide_height).inches
    if abs(w - course["slide"]["width"]) > 0.01 or abs(h - course["slide"]["height"]) > 0.01:
        reds.append(f"畫幅 {w:.2f}×{h:.2f} ≠ 設定 {course['slide']['width']}×{course['slide']['height']}")

    bad_font, small, smallish, placeholder, pagenum = [], [], [], [], []
    runs = 0
    types = {s["num"]: s["type"] for s in built["slides"]}
    for idx, slide in enumerate(prs.slides, 1):
        for p in iter_paras(slide):
            ptxt = p.text.strip()
            if re.fullmatch(r"\d{1,3}", ptxt) and types.get(idx) not in ("stats", "table3", "diagram"):
                pagenum.append(idx)
            if "圖片待補" in ptxt or "▦" in ptxt:
                placeholder.append(idx)
            is_credit = ptxt.startswith(CREDIT_PREFIX)
            for r in p.runs:
                if not r.text.strip():
                    continue
                runs += 1
                rpr = r._r.find(A_NS + "rPr")
                lat = rpr.find(A_NS + "latin") if rpr is not None else None
                ea = rpr.find(A_NS + "ea") if rpr is not None else None
                if (lat is None or lat.get("typeface") != FONT) or (ea is None or ea.get("typeface") != FONT):
                    bad_font.append(idx)
                sz = r.font.size.pt if r.font.size else None
                if sz is None:
                    continue
                if is_credit:
                    if sz < S["credit"]:
                        small.append((idx, r.text[:12], sz))
                elif sz < S["min"]:
                    small.append((idx, r.text[:12], sz))
                elif sz < S["body_warn"]:
                    smallish.append((idx, sz))
    if bad_font:
        reds.append(f"非 {FONT} 的 run：張 {sorted(set(bad_font))[:15]}")
    else:
        print(f"[ok] 字型 {runs} runs 全為 {FONT}")
    if small:
        reds.append(f"字級低於下限（出處 {S['credit']}pt、其他 {S['min']}pt）{len(small)} 個：{small[:10]}")
    else:
        print(f"[ok] 字級：出處 ≥{S['credit']}pt，其餘 ≥{S['min']}pt")
    if smallish:
        warns.append(f"{len(smallish)} 個 run 介於 {S['min']}–{S['body_warn']}pt（張 {sorted(set(i for i, _ in smallish))[:20]}）")
    if placeholder:
        reds.append(f"殘留佔位字：張 {placeholder}")
    if pagenum:
        reds.append(f"疑似靜態頁碼：張 {sorted(set(pagenum))[:10]}")

    over_r, over_w = [], []
    for s in built["slides"]:
        if s["type"] in NO_COUNT:
            continue
        txt = visible_text(dict(s, table=[])) if s["type"] == "table3" else visible_text(s)
        c = len(re.sub(r"\s", "", txt))
        if c > S["char_red"]:
            over_r.append((s["num"], s["id"], c))
        elif c > S["char_warn"]:
            over_w.append((s["num"], s["id"], c))
    if over_r:
        reds.append(f"字元 >{S['char_red']}：{over_r}")
    if over_w:
        warns.append(f"字元 >{S['char_warn']}（建議拆張或刪字）{len(over_w)} 張：{over_w[:12]}")

    content = [s for s in built["slides"] if s["type"] not in NO_COUNT]
    with_img = [s for s in content if any(i["exists"] for i in s["imgs"]) or s["type"] == "diagram"]
    wanted = [s for s in content if not any(i["exists"] for i in s["imgs"]) and s.get("img_wanted")]
    rate = len(with_img) / len(content) if content else 1
    msg = (f"有圖率 {len(with_img)}/{len(content)}＝{rate:.0%}（門檻 {S['img_rate_min']:.0%}；"
           f"待補圖另有 {len(wanted)} 張，全張 {len(with_img)}/{n}＝{len(with_img) / n:.0%}）")
    if rate < S["img_rate_min"]:
        reds.append(msg)
    else:
        print("[ok] " + msg)
    missing = [s["num"] for s in built["slides"] if any(not i["exists"] for i in s["imgs"])]
    if missing:
        warns.append(f"素材池寫了圖但檔案不存在：張 {missing}")

    qs = built["questions"]
    quiz = [s for s in built["slides"] if s["type"] == "quiz"]
    badq = [(s["num"], s["quiz"]) for s in quiz if s["quiz"] not in qs]
    if badq:
        reds.append(f"國考題 key 查無題庫：{badq}")
    elif quiz:
        print(f"[ok] 國考題 {len(quiz)} 題 key 皆有效")
    if quiz and not built["exam"]:
        reds.append("此 deck 設定不放國考題，卻有 quiz 張")
    lo, hi = course.deck(deck).get("per_section", [0, 10 ** 6])
    for a, b, t in built["sections"]:
        k = b - a + 1
        if lo <= k <= hi:
            print(f"[ok] {t}：{k} 張（範圍 {lo}–{hi}）")
        else:
            warns.append(f"{t}：{k} 張，超出 {lo}–{hi}")

    rep_path = os.path.splitext(pptx_path)[0] + "_layout.json"
    if os.path.exists(rep_path):
        rep = json.load(open(rep_path, encoding="utf-8"))
        for num, sid, e in rep.get("errors", []):
            reds.append(f"渲染錯誤 #{num} {sid}: {e}")
        for sid, v in rep["slides"].items():
            for f in v["fits"]:
                if f["overflow"]:
                    (reds if f["need_in"] > f["avail_in"] * 1.15 else warns).append(
                        f"#{v['num']} {sid} {f['what']} 估算需 {f['need_in']}in > 可用 {f['avail_in']}in")
    for m in warns:
        print("[WARN] " + m)
    for m in reds:
        print("[RED] " + m)
    ok = not reds
    print("===== " + ("ALL GREEN" if ok else f"HAS {len(reds)} RED") + " =====")
    return ok


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("--course", required=True)
    ap.add_argument("--pptx", required=True)
    ap.add_argument("--deck", default=None)
    ap.add_argument("--section", type=int, default=None)
    a = ap.parse_args()
    c = load_course(a.course)
    sys.exit(0 if main(c, a.pptx, a.deck or c["deck_order"][0], a.section) else 1)
