# -*- coding: utf-8 -*-
"""任務 B（在教師的 deck 上改版）：從教師自己用 PowerPoint 匯出的 PDF，把每張投影片的「圖區」原樣切成一張圖。

和重繪（LibreOffice）不同：直接用他的 PDF（PowerPoint 算圖、微軟正黑體），所以保留裁切、多圖排版、箭頭、
圖上標籤、SmartArt、群組，跟上課時看到的一樣。

做法：
  1. python-pptx 讀每張的圖片、群組、SmartArt、箭頭等圖形，以及設定檔指定要留的短標籤，取聯集方框（裁到投影片內）。
     畫布外的殘留與極小圖形（< 0.3×0.3 吋；他舊稿常見，指南 AUD-31）一律跳過，不然方框會被拉到投影片邊上。
  2. PDF 那一頁上，中心點不在「要留的標籤／SmartArt／群組」方框裡的文字一律用 redaction 拿掉（圖片與向量圖形不動）。
  3. 方框以 200 dpi 算圖（長邊最多 2000 px）；色數少（示意圖、SmartArt、路徑圖）存 PNG，照片存 JPEG。

設定檔（JSON，可省略）：
  {"labels": {"3": ["一層扁的雨傘細胞", "文字方塊 5"], …},   要跟圖一起留下的短標籤：文字開頭或 shape 名都可以
   "drop":   {"115": ["Rectangle 8"], …}}                     不要的圖形（shape 名）
  他的 deck 文字方塊常同名（"Text Box 4"），用文字開頭比較穩。

用法：python compose_from_pdf.py <in.pptx> <in.pdf> <代號> <輸出資料夾> [--config cfg.json] [--only 3 11 12]
輸出：<輸出資料夾>/<代號>_s{張:03d}.png|jpg、<代號>_compose.json（每張的方框、像素、拿掉幾段字、各部件）
      --only 時只重切那幾張，compose.json 其他張保留。
範例：2026-10 skin（代號 a）、uro-a3（L1 代號 a，L2 用 b）。
"""
import argparse, json, os
import fitz
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU = 914400
DPI = 200
MAX_SIDE = 2000
LABELS, DROP = {}, {}


def is_label(n, sh):
    want = LABELS.get(n, [])
    if sh.name in want:
        return True
    if not sh.has_text_frame:
        return False
    t = sh.text_frame.text.strip()
    return bool(t) and any(t.startswith(p) for p in want)


def off_canvas(sh, SW, SH):
    x0, y0 = sh.left / EMU, sh.top / EMU
    x1, y1 = x0 + sh.width / EMU, y0 + sh.height / EMU
    if x1 <= 0.05 or y1 <= 0.05 or x0 >= SW - 0.05 or y0 >= SH - 0.05:
        return True
    return (x1 - x0) < 0.3 and (y1 - y0) < 0.3


def kind(sh):
    st = sh.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "pic"
    if st == MSO_SHAPE_TYPE.PLACEHOLDER:
        try:
            sh.image
            return "pic"
        except Exception:
            return None
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if "drawingml/2006/diagram" in sh._element.xml:
        return "smartart"
    if st in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.FREEFORM):
        if sh.has_text_frame and sh.text_frame.text.strip():
            return None
        return "shape"
    return None


def main():
    ap = argparse.ArgumentParser(description="從教師 PowerPoint 匯出的 PDF 切出每張的圖區")
    ap.add_argument("pptx"); ap.add_argument("pdf"); ap.add_argument("code"); ap.add_argument("out")
    ap.add_argument("--config", help="JSON：labels（要留的標籤：文字開頭或 shape 名）、drop（不要的 shape 名）")
    ap.add_argument("--only", nargs="*", type=int, default=[], help="只重切這幾張")
    a = ap.parse_args()
    if a.config:
        cfg = json.load(open(a.config, encoding="utf-8"))
        LABELS.update({int(k): list(v) for k, v in (cfg.get("labels") or {}).items()})
        DROP.update({int(k): set(v) for k, v in (cfg.get("drop") or {}).items()})
    src, pdf, code, out = a.pptx, a.pdf, a.code, a.out
    only = set(a.only)
    os.makedirs(out, exist_ok=True)
    prs = Presentation(src)
    SW, SH = prs.slide_width / EMU, prs.slide_height / EMU
    doc = fitz.open(pdf)
    assert len(doc) == len(prs.slides), (len(doc), len(prs.slides))
    meta_path = os.path.join(out, f"{code}_compose.json")
    meta = json.load(open(meta_path, encoding="utf-8")) if (only and os.path.exists(meta_path)) else {}
    for n, slide in enumerate(prs.slides, 1):
        if only and n not in only:
            continue
        parts, keep_text = [], []
        for sh in slide.shapes:
            if sh.name in DROP.get(n, set()) or off_canvas(sh, SW, SH):
                continue
            k = kind(sh)
            if is_label(n, sh):
                k = "label"
            if not k:
                continue
            box = [sh.left / EMU, sh.top / EMU, (sh.left + sh.width) / EMU, (sh.top + sh.height) / EMU]
            parts.append({"kind": k, "name": sh.name, "box_in": [round(v, 3) for v in box]})
            if k in ("label", "smartart", "group"):
                keep_text.append(box)
        if not any(p["kind"] in ("pic", "smartart", "group") for p in parts):
            meta.pop(str(n), None)
            continue
        x0 = max(0, min(p["box_in"][0] for p in parts)); y0 = max(0, min(p["box_in"][1] for p in parts))
        x1 = min(SW, max(p["box_in"][2] for p in parts)); y1 = min(SH, max(p["box_in"][3] for p in parts))
        page = doc[n - 1]
        sx, sy = page.rect.width / SW, page.rect.height / SH
        # 拿掉不屬於圖的文字
        red = 0
        for b in page.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                for sp in ln["spans"]:
                    if not sp["text"].strip():
                        continue
                    r = fitz.Rect(sp["bbox"])
                    cx, cy = (r.x0 + r.x1) / 2 / sx, (r.y0 + r.y1) / 2 / sy
                    if any(k[0] - 0.05 <= cx <= k[2] + 0.05 and k[1] - 0.05 <= cy <= k[3] + 0.05 for k in keep_text):
                        continue
                    page.add_redact_annot(r, fill=False)
                    red += 1
        if red:
            page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE, graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                                  text=fitz.PDF_REDACT_TEXT_REMOVE)
        clip = fitz.Rect(x0 * sx, y0 * sy, x1 * sx, y1 * sy)
        pix = page.get_pixmap(dpi=DPI, clip=clip, alpha=False)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        if max(im.size) > MAX_SIDE:
            f = MAX_SIDE / max(im.size)
            im = im.resize((round(im.width * f), round(im.height * f)), Image.LANCZOS)
        small = im.resize((max(1, im.width // 4), max(1, im.height // 4)))
        cols = small.getcolors(maxcolors=6000)
        png = cols is not None and len(cols) < 2500
        fn = f"{code}_s{n:03d}." + ("png" if png else "jpg")
        for ext in ("png", "jpg"):
            old = os.path.join(out, f"{code}_s{n:03d}.{ext}")
            if os.path.exists(old):
                os.remove(old)
        if png:
            im.save(os.path.join(out, fn), optimize=True)
        else:
            im.save(os.path.join(out, fn), quality=90, optimize=True)
        meta[str(n)] = {"file": fn, "box_in": [round(v, 3) for v in (x0, y0, x1, y1)], "px": [im.width, im.height],
                        "redacted_spans": red, "parts": parts}
    meta = dict(sorted(meta.items(), key=lambda kv: int(kv[0])))
    json.dump(meta, open(meta_path, "w", encoding="utf-8", newline="\n"), ensure_ascii=False, indent=1)
    print(code, len(meta), "images")


if __name__ == "__main__":
    main()
