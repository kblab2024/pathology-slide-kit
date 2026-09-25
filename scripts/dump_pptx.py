# -*- coding: utf-8 -*-
"""把 pptx 逐張傾印成 JSON＋LLM 好讀的純文字，供風格研究與改稿比對。

與 2026-09-07 舊版（ENT 專案時期的傾印腳本，不在 repo）的差異：
- 字級取「實際顯示值」：run → 形狀 lstStyle → 版面配置 placeholder → 母片 placeholder
  → 母片 txStyles（title/body/other）→ 18pt 預設，並乘上 normAutofit 的 fontScale。
- 顏色解析主題色（schemeClr → theme 色票），輸出 hex 與來源（rgb／theme:accent2）。
- 圖片計入群組內、placeholder 內、以圖片填滿的形狀。
- 另出一份 .txt：每張一段，行尾標字級與非預設色，方便 LLM 直接讀。

用法：
  python dump_pptx.py <in.pptx> <out_base>          → out_base.json、out_base.txt
  python dump_pptx.py --batch <list.tsv> <out_dir>  → list.tsv 每行「代號<TAB>pptx 路徑」
"""
import json
import os
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
A = "{%s}" % NS["a"]
P = "{%s}" % NS["p"]


def inch(v):
    return round(Emu(v).inches, 2) if v is not None else None


# ---------------- 主題色 ----------------
def theme_colors(master):
    out = {}
    try:
        th = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(th.blob)
        cs = root.find(".//a:clrScheme", NS)
        for el in cs:
            name = etree.QName(el).localname
            c = el.find("a:srgbClr", NS)
            if c is not None:
                out[name] = c.get("val").upper()
            else:
                s = el.find("a:sysClr", NS)
                if s is not None:
                    out[name] = (s.get("lastClr") or "000000").upper()
    except Exception:
        pass
    # clrMap（tx1→dk1 等）
    try:
        cm = master._element.find(P + "clrMap")
        for k, v in cm.attrib.items():
            if v in out:
                out[k] = out[v]
    except Exception:
        out.setdefault("tx1", out.get("dk1", "000000"))
        out.setdefault("bg1", out.get("lt1", "FFFFFF"))
    return out


def color_of(parent, theme):
    """parent 是 a:rPr 或 a:defRPr；回傳 (hex, 來源) 或 (None, None)。"""
    if parent is None:
        return None, None
    sf = parent.find("a:solidFill", NS)
    if sf is None:
        return None, None
    c = sf.find("a:srgbClr", NS)
    if c is not None:
        return c.get("val").upper(), "rgb"
    c = sf.find("a:schemeClr", NS)
    if c is not None:
        name = c.get("val")
        mods = "".join(f"/{etree.QName(m).localname}{m.get('val')}" for m in c)
        return theme.get(name, "?"), f"theme:{name}{mods}"
    return None, None


# ---------------- 字級繼承 ----------------
def _lvl_def(lst, lvl):
    if lst is None:
        return None
    return lst.find(f"a:lvl{lvl + 1}pPr/a:defRPr", NS)


def _ph_key(sh):
    try:
        ph = sh.placeholder_format
        return ph.type, ph.idx
    except Exception:
        return None


def _find_ph(container, key):
    if key is None:
        return None
    ptype, idx = key
    best = None
    for sh in container.placeholders:
        k = _ph_key(sh)
        if k is None:
            continue
        if k[1] == idx and idx != 0:
            return sh
        if k[0] == ptype or (ptype in (PP_PLACEHOLDER.CENTER_TITLE,) and k[0] == PP_PLACEHOLDER.TITLE):
            best = best or sh
    return best


def _style_kind(key):
    if key is None:
        return "other"
    t = key[0]
    if t in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
        return "title"
    if t in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE):
        return "body"
    return "other"


def inherited(sh, slide, lvl, attr):
    """回傳 (值, 來源)。attr='sz' 或 'color'。"""
    key = _ph_key(sh)
    chain = []
    try:
        chain.append(("shape", sh._element.find(".//p:txBody/a:lstStyle", NS)))
    except Exception:
        pass
    if key is not None:
        lph = _find_ph(slide.slide_layout, key)
        if lph is not None:
            chain.append(("layout", lph._element.find(".//p:txBody/a:lstStyle", NS)))
        mph = _find_ph(slide.slide_layout.slide_master, key)
        if mph is not None:
            chain.append(("master_ph", mph._element.find(".//p:txBody/a:lstStyle", NS)))
    kind = _style_kind(key)
    tx = slide.slide_layout.slide_master._element.find(P + "txStyles")
    if tx is not None:
        tag = {"title": "titleStyle", "body": "bodyStyle", "other": "otherStyle"}[kind]
        chain.append(("master_" + kind, tx.find(P + tag)))
    for src, lst in chain:
        d = _lvl_def(lst, lvl)
        if d is None:
            continue
        if attr == "sz" and d.get("sz"):
            return int(d.get("sz")) / 100.0, src
        if attr == "color":
            yield_c = d.find("a:solidFill", NS)
            if yield_c is not None:
                return d, src
    return (18.0, "default") if attr == "sz" else (None, None)


def autofit_scale(sh):
    try:
        na = sh._element.find(".//a:bodyPr/a:normAutofit", NS)
        if na is not None and na.get("fontScale"):
            return int(na.get("fontScale")) / 100000.0
    except Exception:
        pass
    return 1.0


# ---------------- 形狀 ----------------
def has_pic_fill(sh):
    try:
        return sh._element.find(".//a:blipFill", NS) is not None and sh.shape_type != MSO_SHAPE_TYPE.PICTURE
    except Exception:
        return False


def para_info(p, sh, slide, theme, scale, in_table=False):
    lvl = p.level
    runs = []
    for r in p.runs:
        rpr = r._r.find("a:rPr", NS)
        if rpr is not None and rpr.get("sz"):
            sz, ssrc = int(rpr.get("sz")) / 100.0, "run"
        elif in_table:
            sz, ssrc = 18.0, "table_default"
        else:
            sz, ssrc = inherited(sh, slide, lvl, "sz")
        sz = round(sz * scale, 1)
        col, csrc = color_of(rpr, theme)
        if col is None and not in_table:
            d, dsrc = inherited(sh, slide, lvl, "color")
            if d is not None:
                col, csrc = color_of(d, theme)
                csrc = f"{dsrc}:{csrc}"
        b = r.font.bold
        runs.append({"t": r.text, "sz": sz, "sz_src": ssrc, "b": b, "col": col, "col_src": csrc})
    return {"lvl": lvl, "text": p.text, "runs": runs}


def shape_info(sh, slide, theme, path=""):
    d = {"name": sh.name, "kind": str(sh.shape_type).split(".")[-1].split(" ")[0],
         "x": inch(sh.left), "y": inch(sh.top), "w": inch(sh.width), "h": inch(sh.height)}
    key = _ph_key(sh)
    if key is not None:
        d["ph"] = str(key[0]).split(".")[-1].split(" ")[0]
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or getattr(sh, "image", None) is not None and key is not None:
        try:
            d["img"] = {"ext": sh.image.ext, "bytes": len(sh.image.blob), "px": list(_px(sh.image.blob))}
        except Exception:
            d["img"] = {}
    elif has_pic_fill(sh):
        d["img"] = {"fill": True}
    if sh.has_text_frame and sh.text_frame.text.strip():
        sc = autofit_scale(sh)
        d["autofit"] = sc if sc != 1.0 else None
        d["paras"] = [para_info(p, sh, slide, theme, sc) for p in sh.text_frame.paragraphs if p.text.strip()]
    if getattr(sh, "has_table", False) and sh.has_table:
        rows = []
        for row in sh.table.rows:
            cells = []
            for c in row.cells:
                ps = [para_info(p, sh, slide, theme, 1.0, in_table=True) for p in c.text_frame.paragraphs if p.text.strip()]
                cells.append({"text": c.text, "paras": ps})
            rows.append(cells)
        d["table"] = rows
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        d["children"] = [shape_info(c, slide, theme) for c in sh.shapes]
    if sh._element.tag.endswith("graphicFrame") and "drawingml/2006/diagram" in etree.tostring(sh._element).decode("utf-8", "ignore"):
        d["smartart"] = True
    return d


def _px(blob):
    try:
        from PIL import Image
        import io
        return Image.open(io.BytesIO(blob)).size
    except Exception:
        return (None, None)


def walk(shapes):
    for sh in shapes:
        yield sh
        for c in sh.get("children", []):
            yield from walk([c])


def slide_title(s, shapes):
    if s.shapes.title is not None and s.shapes.title.has_text_frame and s.shapes.title.text.strip():
        return s.shapes.title.text.strip(), "placeholder"
    # 沒有標題 placeholder：取最上方、字最大的文字
    best = None
    for sh in walk(shapes):
        for p in sh.get("paras", []):
            szs = [r["sz"] for r in p["runs"] if r["sz"]]
            if not szs or not p["text"].strip():
                continue
            score = (-(sh["y"] or 0) + max(szs) / 10.0)
            if best is None or score > best[0]:
                best = (score, p["text"].strip())
            break
    return (best[1], "guess") if best else ("", "none")


def dump(path):
    prs = Presentation(path)
    out = {"file": os.path.basename(path), "w": inch(prs.slide_width), "h": inch(prs.slide_height),
           "slides": []}
    for i, s in enumerate(prs.slides, 1):
        theme = theme_colors(s.slide_layout.slide_master)
        shapes = [shape_info(sh, s, theme) for sh in s.shapes]
        title, tsrc = slide_title(s, shapes)
        npic = sum(1 for sh in walk(shapes) if "img" in sh)
        bg = None
        try:
            if s.background.fill.type is not None:
                bg = str(s.background.fill.fore_color.rgb)
        except Exception:
            pass
        out["slides"].append({
            "n": i, "layout": s.slide_layout.name, "title": title, "title_src": tsrc,
            "pics": npic, "smartart": any(sh.get("smartart") for sh in walk(shapes)),
            "bg": bg, "shapes": shapes,
            "notes": s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""})
    return out


# ---------------- LLM 好讀的純文字 ----------------
DEFAULT_COLS = {None, "000000", "1E2A35", "14263B", "3B3838", "404040", "262626", "595959"}


def _fmt_para(p):
    szs = sorted({r["sz"] for r in p["runs"] if r["sz"]})
    sz = "/".join(f"{x:g}" for x in szs) + "pt" if szs else ""
    marks = []
    for r in p["runs"]:
        if r["col"] and r["col"] not in DEFAULT_COLS and r["t"].strip():
            marks.append(f"{r['col']}「{r['t'].strip()}」")
        elif r["b"] and r["t"].strip() and len(p["runs"]) > 1:
            marks.append(f"粗「{r['t'].strip()}」")
    ind = "  " * p["lvl"]
    return f"{ind}{p['text'].strip()}  [{sz}]" + (("  {" + "；".join(marks) + "}") if marks else "")


def to_text(d):
    L = [f"# {d['file']}  ({len(d['slides'])} 張, {d['w']}x{d['h']} in)"]
    for s in d["slides"]:
        flags = []
        if s["pics"]:
            flags.append(f"圖{s['pics']}")
        if s["smartart"]:
            flags.append("SmartArt")
        L.append(f"===== #{s['n']} [{s['layout']}] {' '.join(flags)}")
        for sh in walk(s["shapes"]):
            for p in sh.get("paras", []):
                L.append(_fmt_para(p))
            if "table" in sh:
                for row in sh["table"]:
                    szs = sorted({r["sz"] for c in row for p in c["paras"] for r in p["runs"] if r["sz"]})
                    L.append("| " + " | ".join(c["text"].replace("\n", " ") for c in row) + " |"
                             + (f"  [{'/'.join(f'{x:g}' for x in szs)}pt]" if szs else ""))
            if "img" in sh:
                L.append(f"  [圖 {sh['w']}x{sh['h']}in @({sh['x']},{sh['y']})]")
        if s["notes"].strip():
            L.append("  〔備註〕" + s["notes"].strip().replace("\n", " ⏎ "))
    return "\n".join(L) + "\n"


def run_one(src, base):
    d = dump(src)
    json.dump(d, open(base + ".json", "w", encoding="utf-8", newline="\n"), ensure_ascii=False, indent=1)
    open(base + ".txt", "w", encoding="utf-8", newline="\n").write(to_text(d))
    return d


def main():
    import argparse
    ap = argparse.ArgumentParser(description="把 pptx 逐張傾印成 <out_base>.json 與 <out_base>.txt（風格研究與改稿比對用）")
    ap.add_argument("pptx", nargs="?", help="要傾印的 .pptx")
    ap.add_argument("out_base", nargs="?", help="輸出路徑（不含副檔名），例 <materials>/corpus/decks/<代號>")
    ap.add_argument("--batch", nargs=2, metavar=("LIST_TSV", "OUT_DIR"),
                    help="整批傾印：LIST_TSV 每行「代號<TAB>pptx 路徑」，輸出到 OUT_DIR/<代號>.json|txt")
    a = ap.parse_args()
    if a.batch:
        lst, outdir = a.batch
        os.makedirs(outdir, exist_ok=True)
        for ln in open(lst, encoding="utf-8"):
            if not ln.strip() or ln.startswith("#"):
                continue
            code, src = ln.rstrip("\r\n").split("\t", 1)
            d = run_one(src, os.path.join(outdir, code))
            print(code, len(d["slides"]), d["w"], "x", d["h"])
        return 0
    if not (a.pptx and a.out_base):
        ap.error("需要 <pptx> <out_base>，或 --batch <list.tsv> <out_dir>")
    d = run_one(a.pptx, a.out_base)
    print(os.path.basename(a.pptx), len(d["slides"]))
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
