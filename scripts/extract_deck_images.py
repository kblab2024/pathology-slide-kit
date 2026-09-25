# -*- coding: utf-8 -*-
"""從既有 pptx 抽出所有圖片（含群組內、placeholder 內），連同張號、同張文字與圖說一起建清單。

用途：重用教師舊投影片裡的 ExpertPath、期刊圖、Robbins 舊版截圖。
輸出：<out>/<代號>_s{張:03d}_{序}.{副檔名}、<out>/<代號>_images.json、並在 <out>/圖片清單.md 追加一節。
同一張圖（內容 sha1 相同）只存一次，後出現者記為 dup_of。

用法：python extract_deck_images.py <in.pptx> <out_dir> --code <代號> [--slides 1-47] [--author 教師]
  --slides  只抽這些張（例：2024 免疫一只有 1–47 張是教師自製）
  --author  寫進清單的作者欄（非教師本人的圖，重用時封面要致謝）
"""
import argparse
import hashlib
import json
import os
import re
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

CAP_PAT = re.compile(r"(圖源|圖片來源|資料來源|來源|Robbins|ROBBINS|Fig\.|http|www\.|doi|PMID|et al|©|Courtesy|ExpertPath)", re.I)


def walk(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(sh.shapes)
        else:
            yield sh


def parse_range(s, n):
    if not s:
        return set(range(1, n + 1))
    out = set()
    for part in s.split(","):
        a, _, b = part.partition("-")
        out |= set(range(int(a), int(b or a) + 1))
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("out")
    ap.add_argument("--code", required=True)
    ap.add_argument("--slides")
    ap.add_argument("--author", default="教師")
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    prs = Presentation(a.pptx)
    keep = parse_range(a.slides, len(prs.slides))
    seen, recs = {}, []
    for i, s in enumerate(prs.slides, 1):
        if i not in keep:
            continue
        texts = [sh.text_frame.text.strip() for sh in walk(s.shapes) if sh.has_text_frame and sh.text_frame.text.strip()]
        caps = [t for t in texts if CAP_PAT.search(t)]
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        k = 0
        for sh in walk(s.shapes):
            try:
                img = sh.image
            except Exception:
                continue
            k += 1
            blob = img.blob
            h = hashlib.sha1(blob).hexdigest()[:16]
            ext = img.ext if img.ext != "jpeg" else "jpg"
            name = f"{a.code}_s{i:03d}_{k}.{ext}"
            rec = {"file": name, "slide": i, "k": k, "sha1": h, "ext": ext,
                   "pos_in": [round(Emu(sh.left).inches, 2), round(Emu(sh.top).inches, 2),
                              round(Emu(sh.width).inches, 2), round(Emu(sh.height).inches, 2)],
                   "slide_text": " ／ ".join(texts)[:400], "captions": caps[:4],
                   "notes": notes[:300], "author": a.author, "deck": os.path.basename(a.pptx)}
            try:
                from PIL import Image
                import io
                rec["px"] = list(Image.open(io.BytesIO(blob)).size)
            except Exception:
                rec["px"] = None
            if h in seen:
                rec["dup_of"] = seen[h]
            else:
                open(os.path.join(a.out, name), "wb").write(blob)
                seen[h] = name
            recs.append(rec)
    json.dump(recs, open(os.path.join(a.out, f"{a.code}_images.json"), "w", encoding="utf-8", newline="\n"),
              ensure_ascii=False, indent=1)
    uniq = [r for r in recs if "dup_of" not in r]
    L = [f"\n## {a.code}（{os.path.basename(a.pptx)}；作者 {a.author}；張 {a.slides or '全部'}）", "",
         "| 檔案 | 張 | 像素 | 同張文字（節錄） | 圖說／出處 |", "|---|---|---|---|---|"]
    for r in uniq:
        px = f"{r['px'][0]}×{r['px'][1]}" if r["px"] else ""
        L.append(f"| {r['file']} | {r['slide']} | {px} | {r['slide_text'][:60].replace('|', '／')} | "
                 f"{' ／ '.join(r['captions'])[:80].replace('|', '／')} |")
    open(os.path.join(a.out, "圖片清單.md"), "a", encoding="utf-8", newline="\n").write("\n".join(L) + "\n")
    print(f"{a.code}: {len(recs)} 張圖（去重後 {len(uniq)}）")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
