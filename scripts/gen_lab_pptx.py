# -*- coding: utf-8 -*-
"""實驗講義 PPTX：讀 deck.json，照教師 2026 自製實驗講義（〈病理學實驗_A1_發炎〉）的版型產 16:9 PPTX。

用法：python gen_lab_pptx.py <deck.json> --out <輸出.pptx> [--report <報告.json>] [--no-notes] [--bg-pt 24]
                              [--img-max-side 1600] [--jpeg-quality 82]
      python gen_lab_pptx.py --fit <deck.json 或節檔.json> [--bg-pt 24] [--bullets]
        只量字、不產檔：逐張印出需要的高度對可用高度（撰寫代理自查用；量法和正式產檔同一套）

deck.json（實驗講義工作流的輸出）：
  {"cover": "病理學實驗 骨與關節",
   "sections": [{"section": 1,
                 "divider": {"label": "Section 1", "title_lines": ["…", "(…)"], "slides_line": "MC01"},
                 "slides": [{"id": "MC01-1", "kind": "shot", "title": "…", "text": ["句子1", "句子2"],
                             "text_side": "right", "img": {"slot": "left", "mag": "低倍(2x)", "capture": "…"},
                             "sources": ["…"]}]}]}

每張內容張的 kind（沒寫就是 shot；2026-10-15 教師：「實驗講義不可以只有截圖，要有背景知識跟形態學示意圖」）：
  shot  截圖張：有 img.path 就放那張截圖（wsi_crop.py 從掃描檔切、annotate_shot.py 加註記；2026-10-21 起）；
        沒有就放不含任何字的淡灰空圖（形狀名「截圖待補 <id>」），教師用 PowerPoint「變更圖片」換截圖；
        備註寫倍率與視野
  bg    背景知識張：沒有圖，標題下全寬條列（--bg-pt，預設 24pt；「  - 」開頭是第二層）；
        可加 "table": [[表頭…], [列…], …]（第一列是表頭、第一欄粗體），表格放不下就接續一張同標題、重複表頭
  gx    大體／影像張：img.path 的圖等比放進圖框（不拉伸、置中），圖下 14pt 出處行 img.credit；
        沒有圖檔就放淡灰空圖（形狀名「圖待補 <id>」）
  sc    形態學示意圖張：img.path 有圖就照 gx 放；沒有就放淡灰空圖（形狀名「示意圖待換 <id>」，
        之後由 lab_swap_images.py 換成加好標籤的圖）
  img.path 相對 deck.json 所在的資料夾。gx、sc、shot 的說明句都放圖框另一邊（20.2pt）。
deck.json 最上層 "panel": true：圖框與說明句底下各墊一塊淺灰底（EAE9E6，6.14×5.32 吋），照教師發炎講義的原樣；
  沒寫就不墊（骨與關節實驗講義是不墊的版本）。"bullets": true：說明句每段前加「•」（懸掛縮排 0.27 吋、段距 14pt），
  也是發炎講義的原樣。--fit 量 deck.json 時照檔裡的 bullets 量；量單張或節檔時用 --bullets。

版型（幾何取自教師的發炎講義，單位吋，投影片 13.33×7.5）：
  封面      一行 54pt，置中
  分隔張    藍底 143F9C、白字：「Section N」13.5pt，主題 45pt（一到兩行），玻片行 42pt；左緣 1.24、寬 10.86
  內容張    標題一行 30pt 置中（y 0.44）；圖框 6.14×5.32 在 (0.52,1.66)（左）或 (6.67,1.66)（右），
            另一邊放 20.2pt 說明句（寬 5.58、垂直置中）
  背景張    標題同上；條列區 (0.52,1.35) 起、寬 12.29、高 5.75
  字型一律微軟正黑體（common.FONT）。
截圖要拍的倍率與視野、資料來源寫進備註欄（--no-notes 不寫）；gx、sc、bg 的備註只寫來源。

量字：沿用 gen_pptx 的字型量測（Windows 微軟正黑體；Mac 用 PowerPoint 內建的 MSJH；雲端 Noto Sans CJK TC）。
報出：
  overflow  放不下（說明句或背景條列總高超過可用高度、表格單列放不下、分隔張主題或玻片行超過一行）
  rule      違反版型規則但仍放得下（標題超過一行、說明句超過 6 行）
標題裡的「\n」是手動換行（例：「DS22 瀰漫性大 B 細胞淋巴瘤\n(DIFFUSE LARGE B-CELL LYMPHOMA)低倍」），兩行仍算 rule。
有 overflow 時結束碼 1。報告另寫 ids（id → 張號；背景張接續時寫第一張）與各種圖位的格數。
"""
import argparse
import io
import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import bullet_level  # noqa: E402
from gen_pptx import TABLE_LINE_K, _font, em, line_h, measure_font_info, n_lines  # noqa: E402

SW, SH = 13.33, 7.5
BLUE = "143F9C"
INK = "000000"
WHITE = "FFFFFF"
GREYTXT = "595959"
TINT = "EEF2FA"
GREY = (230, 230, 230)
PANEL = "EAE9E6"   # 教師發炎講義圖框與說明句底下的淺灰底
BUL_IND, BUL_SPC = 0.266, 14   # 發炎講義說明句的「•」懸掛縮排（吋）與段距（pt）
STYLE = {"bullets": False}
SZ = {"cover": 54, "label": 13.5, "theme": 45, "glass": 42, "title": 30, "body": 20.2, "credit": 14}
IMG_W, IMG_H, IMG_Y = 6.14, 5.32, 1.66
IMG_X = {"left": 0.52, "right": 6.67}
TXT_W = 5.58
TXT_X = {"right": 7.2, "left": 0.55}   # 說明句在右（圖在左）／在左（圖在右）
DIV_X, DIV_W = 1.24, 10.86
TITLE_Y, TITLE_W = 0.44, 12.93
BODY_MAX_LINES = 6
SPACE_AFTER = 6  # pt
BG_X, BG_Y, BG_W, BG_H = 0.52, 1.35, SW - 2 * 0.52, 5.75
BG_LINE, BG_SPACE = 1.1, 10   # 背景條列的行距倍數、段後 pt
BG_IND = 0.36                 # 條列第一層的懸掛縮排（吋）；每多一層加 0.40
TBL_GAP = 0.15
IMG_MAX_SIDE = 1600
KINDS = ("shot", "bg", "gx", "sc")
SLOT_NAME = {"shot": "截圖待補", "gx": "圖待補", "sc": "示意圖待換"}


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf


def add_para(tf, text, size, color, align=PP_ALIGN.LEFT, first=False, space_after=0, line=1.0, bold=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    if space_after:
        p.space_after = Pt(space_after)
    r = p.add_run()
    r.text = text
    _font(r, size, color, bold)
    r.font.color.rgb = RGBColor.from_string(color)
    return p


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(color)


def add_panel(slide, x, y, w, h):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string(PANEL)
    sp.line.fill.background()
    sp.shadow.inherit = False
    st = sp._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if st is not None:                       # 拿掉主題樣式（effectRef 會帶陰影；LibreOffice 不管 effectLst）
        sp._element.remove(st)
    sp.name = "淺灰底"
    return sp


def grey_png():
    """不含任何字的淡灰空圖（像素比例同教師講義的截圖 590×511）。"""
    buf = io.BytesIO()
    Image.new("RGB", (1180, 1022), GREY).save(buf, "PNG")
    return buf.getvalue()


class Report:
    def __init__(self):
        self.items = []

    def add(self, kind, where, what, need, avail, text):
        self.items.append({"kind": kind, "slide": where, "what": what, "need": need, "avail": avail, "text": text})


# ---------------- 量字（產檔與 --fit 共用） ----------------
def kind_of(c):
    k = c.get("kind") or "shot"
    return k if k in KINDS else "shot"


def slot_side(c):
    img = c.get("img") or {}
    side = img.get("slot") or ("left" if c.get("text_side", "right") == "right" else "right")
    return side if side in IMG_X else "left"


def side_text(c):
    return [t for t in (c.get("text") or []) if str(t).strip()]


def side_need(text):
    w = TXT_W - (BUL_IND if STYLE["bullets"] else 0)
    gap = BUL_SPC if STYLE["bullets"] else SPACE_AFTER
    lines = sum(n_lines(t, SZ["body"], w) for t in text)
    need = lines * line_h(SZ["body"], 1.0) + gap / 72.0 * max(0, len(text) - 1)
    return lines, need


def title_lines(title):
    """標題行數；標題裡的「\n」是手動換行（例：長的英文全名自己成一行，不讓「低倍」單獨掉到第二行）。"""
    return sum(n_lines(seg, SZ["title"], TITLE_W) for seg in str(title).split("\n"))


def check_title(rep, where, title):
    tk = title_lines(title)
    if tk > 1:
        rep.add("rule" if tk == 2 else "overflow", where, "標題超過一行", tk, 1, title)
    return tk


def check_side(rep, where, text):
    lines, need = side_need(text)
    if need > IMG_H + 0.05:
        rep.add("overflow", where, "說明句放不下", round(need, 2), IMG_H, " / ".join(text))
    elif lines > BODY_MAX_LINES:
        rep.add("rule", where, f"說明句超過 {BODY_MAX_LINES} 行", lines, BODY_MAX_LINES, " / ".join(text))
    return lines, need


def bullets_need(bullets, size):
    need = 0.0
    for b in bullets:
        lvl, t = bullet_level(b)
        need += n_lines(t, size, BG_W - (BG_IND + 0.40 * lvl)) * line_h(size, BG_LINE) + BG_SPACE / 72.0
    return need


def table_layout(rows, size):
    """欄寬照字數比例（同 gen_pptx.L_table3）；列高＝最多行的那格 × 行高 × TABLE_LINE_K ＋ 0.14。"""
    ncol = max(len(r) for r in rows)
    rows = [[str(x) for x in (list(r) + [""] * ncol)[:ncol]] for r in rows]
    wr = [max(em(r[j]) for r in rows) for j in range(ncol)]
    tot = sum(max(4, x) for x in wr)
    widths = [max(1.3, BG_W * max(4, x) / tot) for x in wr]
    k = BG_W / sum(widths)
    widths = [x * k for x in widths]
    row_h = [max(n_lines(r[j], size, widths[j] - 0.2, bold=(i == 0 or j == 0)) for j in range(ncol))
             * line_h(size, 1.1) * TABLE_LINE_K + 0.14 for i, r in enumerate(rows)]
    return rows, widths, row_h


def plan_bg(c, size, rep, where):
    """背景張分頁：第一頁放條列（＋放得下的表格列），表格剩下的列接續成同標題的下一張（重複表頭）。
    回傳 [{"bullets": [...], "rows": [列索引…], "need": 高度}]，rows 不含表頭（表頭每頁都放）。"""
    bullets = side_text(c)
    tbl = [r for r in (c.get("table") or []) if r]
    need_b = bullets_need(bullets, size) if bullets else 0.0
    if need_b > BG_H + 0.05:
        rep.add("overflow", where, "背景條列放不下", round(need_b, 2), BG_H, " / ".join(bullets))
    pages = [{"bullets": bullets, "rows": [], "need": need_b}]
    if len(tbl) < 2:
        return pages, None
    rows, widths, row_h = table_layout(tbl, size)
    head = row_h[0]
    cur = pages[0]
    used = need_b + (TBL_GAP if bullets else 0.0)   # 這一頁表格之前已用掉的高度
    for i in range(1, len(rows)):
        h = row_h[i]
        if head + h > BG_H + 0.05:
            rep.add("overflow", where, f"表格第 {i} 列一頁放不下", round(head + h, 2), BG_H, " / ".join(rows[i]))
        add_h = h + (0.0 if cur["rows"] else head)          # 這頁第一列要連表頭一起放
        empty = not cur["rows"] and not cur["bullets"]       # 空白頁：再高也先放（放不下已報 overflow）
        if used + add_h <= BG_H + 0.05 or empty:
            cur["rows"].append(i)
            used += add_h
            cur["need"] = used
        else:
            cur = {"bullets": [], "rows": [i], "need": head + h}
            pages.append(cur)
            used = head + h
    return pages, (rows, widths, row_h)


# ---------------- 產檔 ----------------
def add_title(s, title):
    tk = title_lines(title)
    tf = textbox(s, (SW - TITLE_W) / 2, TITLE_Y, TITLE_W, 0.47 if tk <= 1 else 0.93)
    segs = str(title).split("\n")
    p = add_para(tf, segs[0], SZ["title"], INK, PP_ALIGN.CENTER, first=True)
    for seg in segs[1:]:
        p.add_line_break()
        r = p.add_run()
        r.text = seg
        _font(r, SZ["title"], INK, False)
        r.font.color.rgb = RGBColor.from_string(INK)


def add_side_text(s, side, text):
    tside = "left" if side == "right" else "right"
    tf = textbox(s, TXT_X[tside], IMG_Y, TXT_W, IMG_H, MSO_ANCHOR.MIDDLE)
    bul = STYLE["bullets"]
    for i, t in enumerate(text):
        p = add_para(tf, t, SZ["body"], INK, first=(i == 0),
                     space_after=0 if bul else (SPACE_AFTER if i < len(text) - 1 else 0))
        if bul:
            from lxml import etree
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(BUL_IND * 914400)))
            pPr.set("indent", str(int(-BUL_IND * 914400)))
            if i > 0:
                p.space_before = Pt(BUL_SPC)
            ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
            for tag in ("buSzPct", "buChar"):
                for old in pPr.findall(f"{{{ns}}}{tag}"):
                    pPr.remove(old)
            bs = etree.SubElement(pPr, f"{{{ns}}}buSzPct")
            bs.set("val", "100000")
            bc = etree.SubElement(pPr, f"{{{ns}}}buChar")
            bc.set("char", "•")


IMG_OPT = {"max_side": IMG_MAX_SIDE, "quality": None}   # --img-max-side、--jpeg-quality


def img_bytes(path):
    """圖檔讀進記憶體；長邊超過 max_side 就縮（照片存 JPEG，有透明或 PNG 示意圖存 PNG）。
    給了 --jpeg-quality：沒有透明的圖一律重存 JPEG（截圖多的實驗講義用，150 張從 100 MB 壓到約 35 MB）。"""
    im = Image.open(path)
    im.load()
    w, h = im.size
    ms, q = IMG_OPT["max_side"], IMG_OPT["quality"]
    alpha = im.mode in ("RGBA", "LA", "P") and (im.mode != "P" or "transparency" in im.info)
    if max(w, h) <= ms and (q is None or alpha):
        return open(path, "rb").read(), (w, h)
    if max(w, h) > ms:
        r = ms / max(w, h)
        im = im.resize((max(1, int(w * r)), max(1, int(h * r))), Image.LANCZOS)
    buf = io.BytesIO()
    if alpha or (q is None and path.lower().endswith(".png")):
        im.save(buf, "PNG", optimize=True)
    else:
        im.convert("RGB").save(buf, "JPEG", quality=q or 88)
    return buf.getvalue(), im.size


def add_real_image(s, path, side, credit_text, cid):
    data, (iw, ih) = img_bytes(path)
    x, y, w, h = IMG_X[side], IMG_Y, IMG_W, IMG_H
    t = ""
    if credit_text:
        t = credit_text if credit_text.startswith(("圖", "資料來源", "示意圖")) else "圖：" + credit_text
        k = n_lines(t, SZ["credit"], w)
        h = min(h, SH - 0.08 - 0.25 * k - 0.04 - y)   # 出處行放圖下，不壓圖（2026-10-16 實排：兩行出處壓到直式圖的下緣）
    if iw / ih > w / h:
        dw, dh = w, w * ih / iw
    else:
        dw, dh = h * iw / ih, h
    dx, dy = x + (w - dw) / 2, y + (h - dh) / 2
    pic = s.shapes.add_picture(io.BytesIO(data), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
    pic.name = f"圖 {cid}"
    if t:
        tf = textbox(s, x, dy + dh + 0.04, w, 0.25 * k)
        add_para(tf, t, SZ["credit"], GREYTXT, first=True)
    return pic


def add_bg_page(s, page, layout, size):
    y = BG_Y
    if page["bullets"]:
        tf = textbox(s, BG_X, BG_Y, BG_W, BG_H)
        for i, b in enumerate(page["bullets"]):
            lvl, t = bullet_level(b)
            ind = BG_IND + 0.40 * lvl
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = BG_LINE
            p.space_after = Pt(BG_SPACE)
            pPr = p._p.get_or_add_pPr()
            pPr.set("hangingPunct", "0")
            pPr.set("marL", str(int(ind * 914400)))
            pPr.set("indent", str(int(-0.32 * 914400)))
            r = p.add_run()
            r.text = "• " if lvl == 0 else "– "
            _font(r, size, BLUE if lvl == 0 else GREYTXT, True)
            r = p.add_run()
            r.text = t
            _font(r, size, INK, False)
        y = BG_Y + bullets_need(page["bullets"], size) + TBL_GAP
    if page["rows"] and layout:
        rows, widths, row_h = layout
        idx = [0] + page["rows"]
        hs = [row_h[i] for i in idx]
        t = s.shapes.add_table(len(idx), len(widths), Inches(BG_X), Inches(y), Inches(BG_W), Inches(sum(hs))).table
        t.first_row = False
        t.horz_banding = False
        for j, wj in enumerate(widths):
            t.columns[j].width = Inches(wj)
        for k, (i, hh) in enumerate(zip(idx, hs)):
            t.rows[k].height = Inches(max(0.3, hh))
            for j, txt in enumerate(rows[i]):
                cell = t.cell(k, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor.from_string(BLUE if i == 0 else (WHITE if k % 2 else TINT))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left, cell.margin_right = Inches(0.10), Inches(0.08)
                cell.margin_top, cell.margin_bottom = Inches(0.04), Inches(0.04)
                tf = cell.text_frame
                tf.word_wrap = True
                p = add_para(tf, txt, size, WHITE if i == 0 else INK, first=True, line=1.1, bold=(i == 0 or j == 0))
                p._p.get_or_add_pPr().set("hangingPunct", "0")   # 全形標點不懸到格外（2026-10-16 實排）


def build(deck, out, notes=True, base_dir=".", bg_pt=24):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
    blank = prs.slide_layouts[6]
    png = grey_png()
    rep = Report()
    ids = {}
    slots = {"shot": 0, "shot_image": 0, "gx_image": 0, "gx_placeholder": 0, "sc_image": 0, "sc_placeholder": 0, "bg": 0}
    n = 0

    # 封面
    s = prs.slides.add_slide(blank)
    n += 1
    cover = deck.get("cover", "")
    tf = textbox(s, 0.5, 3.04, SW - 1.0, 0.9, MSO_ANCHOR.MIDDLE)
    add_para(tf, cover, SZ["cover"], INK, PP_ALIGN.CENTER, first=True)
    k = n_lines(cover, SZ["cover"], SW - 1.0)
    if k > 1:
        rep.add("overflow", f"#{n} 封面", "封面超過一行", k, 1, cover)

    shots = 0
    panel = bool(deck.get("panel"))
    STYLE["bullets"] = bool(deck.get("bullets"))
    for sec in deck.get("sections", []):
        d = sec.get("divider", {})
        s = prs.slides.add_slide(blank)
        n += 1
        set_bg(s, BLUE)
        where = f"#{n} {d.get('label', '')}"
        tf = textbox(s, DIV_X, 2.52, DIV_W, 0.24)
        add_para(tf, d.get("label", ""), SZ["label"], WHITE, first=True)
        y = 2.91
        for line in d.get("title_lines", []):
            tf = textbox(s, DIV_X, y, DIV_W, 0.66)
            add_para(tf, line, SZ["theme"], WHITE, first=True)
            k = n_lines(line, SZ["theme"], DIV_W)
            if k > 1:
                rep.add("overflow", where, "分隔張主題超過一行", k, 1, line)
            y += 0.75
        if len(d.get("title_lines", [])) > 2:
            rep.add("rule", where, "分隔張主題超過兩行", len(d["title_lines"]), 2, " / ".join(d["title_lines"]))
        glass = d.get("slides_line", "")
        tf = textbox(s, DIV_X, y - 0.07, DIV_W, 0.78)
        add_para(tf, glass, SZ["glass"], WHITE, first=True)
        k = n_lines(glass, SZ["glass"], DIV_W)
        if k > 1:
            rep.add("overflow", where, "分隔張玻片行超過一行", k, 1, glass)

        for c in sec.get("slides", []):
            kind = kind_of(c)
            cid = c.get("id", "")
            title = c.get("title", "")
            src_note = ("來源：" + "；".join(c["sources"])) if c.get("sources") else ""

            if kind == "bg":
                pages, layout = plan_bg(c, bg_pt, rep, f"#{n + 1} {cid}")
                for pi, page in enumerate(pages):
                    s = prs.slides.add_slide(blank)
                    n += 1
                    if pi == 0:
                        ids[cid] = n
                        check_title(rep, f"#{n} {cid}", title)
                    add_title(s, title)
                    add_bg_page(s, page, layout, bg_pt)
                    slots["bg"] += 1
                    if notes and src_note:
                        s.notes_slide.notes_text_frame.text = src_note
                continue

            s = prs.slides.add_slide(blank)
            n += 1
            ids[cid] = n
            where = f"#{n} {cid}"
            check_title(rep, where, title)
            add_title(s, title)

            img = c.get("img") or {}
            side = slot_side(c)
            path = img.get("path")
            apath = os.path.normpath(os.path.join(base_dir, path)) if path else None
            if panel:
                add_panel(s, IMG_X[side], IMG_Y, IMG_W, IMG_H)
                add_panel(s, IMG_X["right" if side == "left" else "left"], IMG_Y, IMG_W, IMG_H)
            if apath and os.path.exists(apath):
                add_real_image(s, apath, side, img.get("credit", ""), cid)
                if kind == "shot":
                    shots += 1
                    slots["shot_image"] += 1
                else:
                    slots[kind + "_image"] += 1
            else:
                if path:
                    rep.add("rule", where, "圖檔不存在，先放淡灰空圖", 0, 0, path)
                pic = s.shapes.add_picture(io.BytesIO(png), Inches(IMG_X[side]), Inches(IMG_Y), Inches(IMG_W), Inches(IMG_H))
                pic.name = f"{SLOT_NAME[kind]} {cid}"
                if kind == "shot":
                    shots += 1
                    slots["shot"] += 1
                else:
                    slots[kind + "_placeholder"] += 1

            text = side_text(c)
            add_side_text(s, side, text)
            check_side(rep, where, text)

            if notes:
                parts = []
                if kind == "shot" and img:
                    parts.append(f"截圖：{img.get('mag', '')}｜{img.get('capture', '')}")
                if src_note:
                    parts.append(src_note)
                if parts:
                    s.notes_slide.notes_text_frame.text = "\n".join(parts)

    prs.save(out)
    return n, shots, rep, ids, slots


# ---------------- --fit：只量字 ----------------
def fit_slides(obj):
    if isinstance(obj, list):
        return obj
    if "sections" in obj:
        return [c for sec in obj["sections"] for c in sec.get("slides", [])]
    if "slides" in obj:
        return obj["slides"]
    return [obj]


def fit(path, bg_pt, bullets=False):
    obj = json.load(open(path, encoding="utf-8"))
    STYLE["bullets"] = bool(obj.get("bullets")) if isinstance(obj, dict) and "bullets" in obj else bullets
    rep = Report()
    print(f"量字字型：{measure_font_info()}；背景張 {bg_pt}pt、說明句 {SZ['body']}pt")
    print("id｜kind｜標題行數｜需要高度/可用（吋）｜行數")
    for c in fit_slides(obj):
        kind = kind_of(c)
        cid = c.get("id", "?")
        where = cid
        tk = check_title(rep, where, c.get("title", ""))
        if kind == "bg":
            pages, layout = plan_bg(c, bg_pt, rep, where)
            for pi, p in enumerate(pages):
                print(f"{cid}{'' if pi == 0 else f'（接續 {pi + 1}）'}｜bg｜{tk}｜{p['need']:.2f}/{BG_H}｜"
                      f"條列 {len(p['bullets'])} 條、表格 {len(p['rows'])} 列")
            if len(pages) > 1:
                rep.add("rule", where, "背景張接續成多張", len(pages), 1, c.get("title", ""))
        else:
            lines, need = check_side(rep, where, side_text(c))
            print(f"{cid}｜{kind}｜{tk}｜{need:.2f}/{IMG_H}｜{lines} 行（上限 {BODY_MAX_LINES}）")
    ov = [i for i in rep.items if i["kind"] == "overflow"]
    for i in rep.items:
        print(f"  [{i['kind']}] {i['slide']} {i['what']}（{i['need']} > {i['avail']}）：{i['text'][:60]}")
    print(f"放不下 {len(ov)}、版型提醒 {len(rep.items) - len(ov)}")
    return 1 if ov else 0


def main():
    ap = argparse.ArgumentParser(description="實驗講義 deck.json → PPTX（教師 2026 實驗講義版型）")
    ap.add_argument("deck", nargs="?")
    ap.add_argument("--out")
    ap.add_argument("--report", help="放不下與版型問題寫成 JSON（預設：<輸出>_layout.json）")
    ap.add_argument("--no-notes", action="store_true", help="備註欄不寫截圖說明與來源")
    ap.add_argument("--bg-pt", type=float, default=24, help="背景知識張的條列與表格字級（預設 24）")
    ap.add_argument("--fit", metavar="JSON", help="只量字、不產檔（deck.json、節檔或張的清單）")
    ap.add_argument("--bullets", action="store_true", help="--fit 量單張或節檔時照「•」懸掛縮排量（deck.json 以檔裡的 bullets 為準）")
    ap.add_argument("--img-max-side", type=int, default=IMG_MAX_SIDE, help=f"放進 PPTX 的圖長邊上限 px（預設 {IMG_MAX_SIDE}）")
    ap.add_argument("--jpeg-quality", type=int, default=None, help="給了就把沒有透明的圖都重存成這個品質的 JPEG（例 82）")
    a = ap.parse_args()
    if a.fit:
        return fit(a.fit, a.bg_pt, a.bullets)
    if not a.deck or not a.out:
        ap.error("要給 deck.json 與 --out（或用 --fit 只量字）")
    deck = json.load(open(a.deck, encoding="utf-8"))
    IMG_OPT.update(max_side=a.img_max_side, quality=a.jpeg_quality)
    print("量字字型：" + measure_font_info())
    if os.path.exists(a.out):
        print(f"[注意] 覆寫既有檔：{a.out}")
    base = os.path.dirname(os.path.abspath(a.deck))
    n, shots, rep, ids, slots = build(deck, a.out, notes=not a.no_notes, base_dir=base, bg_pt=a.bg_pt)
    rp = a.report or os.path.splitext(a.out)[0] + "_layout.json"
    with open(rp, "w", encoding="utf-8") as f:
        json.dump({"measure_font": measure_font_info(), "slides": n, "image_slots": shots, "slots": slots,
                   "ids": ids, "issues": rep.items}, f, ensure_ascii=False, indent=1)
    print(f"OK {n} 張（截圖圖框 {shots} 格；{slots}）-> {a.out}")
    ov = [i for i in rep.items if i["kind"] == "overflow"]
    for i in rep.items:
        print(f"  [{i['kind']}] {i['slide']} {i['what']}（{i['need']} > {i['avail']}）：{i['text'][:60]}")
    print(f"放不下 {len(ov)}、版型提醒 {len(rep.items) - len(ov)}；報告 {rp}")
    return 1 if ov else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
