# -*- coding: utf-8 -*-
"""把素材池＋選單算成 PPTX（講稿進備註欄），並自動跑 check_deck。

用法：python gen_pptx.py --course immune [--deck dent|all] [--date YYYYMMDD] [--section N] [--name 檔名] [--no-check]
  --section N  只輸出第 N 節（教師有時把每節存成獨立檔）；檔名自動加「_第N節」，不會蓋掉全套檔
  --name       輸出檔名（不含 .pptx）；預設用 course.yaml 的 decks.<deck>.filename。
               多套 deck（--deck all 且 deck_order 不只一套）時自動加「_<deck>」，各套不會互相覆蓋
  寫檔前若同名 PPTX 已存在會印出提醒（老師改過的檔要先回灌，見 SOP 第 11b 步）

版型（pool 的 <!-- type: -->）：
  title 封面｜divider 節或段落分隔｜left 條列＋右圖（無圖則全寬條列）｜flash 標題＋說明句＋大圖
  big 單句大字（可附圖）｜pair 兩圖並排（IMG 可帶「標籤：HE」）｜robbins 整張教科書圖
  table3 表格（超過 table_max_rows 列自動拆張）｜quiz 國考題（course.yaml decks.<deck>.quiz_style：
  乙＝年次當標題、正解選項整段強調色（預設）；甲＝1 列 3 欄表、首欄正解字母粗體）
  stats 大數字｜diagram 框架圖（course.yaml 的 diagrams）
字級全部來自 course.yaml 的 sizes；子層條列與母層同字級（教師 2026 做法）。
另輸出 {檔名}_layout.json：每張估算的文字行數與是否可能溢出，check_deck 會讀它。
"""
import argparse
import hashlib
import json
import os
import sys
import tempfile
import re
import unicodedata

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (DARK_TEXT, FONT, build_deck, bullet_level, deck_filename, load_course,  # noqa: E402
                    notes_map, parse_markup, parse_pool, parse_script, strip_markup)

Image.MAX_IMAGE_PIXELS = None
_TMP = os.path.join(tempfile.gettempdir(), "_deck_imgtmp")
os.makedirs(_TMP, exist_ok=True)
_PREP = {}
CN_NUM = {1: "一", 2: "二", 3: "三", 4: "四"}
NL = chr(10)


class G:
    """渲染期全域設定（由 setup(course) 填入）。"""
    SW = SH = M = 0.0
    S = {}
    D = {}
    EMPH = {}
    course = None
    layout = {}


def setup(course):
    G.course = course
    G.SW, G.SH = course["slide"]["width"], course["slide"]["height"]
    G.M = course["slide"].get("margin", 0.5)
    G.S = course.sizes
    G.D = course["design"]
    G.EMPH = course["emphasis"]
    G.BG = course["slide"].get("background", G.D["snow"])


def C(h):
    return RGBColor.from_string(h)


# ---------------- 圖片 ----------------
def prep_image(path, max_side=1800):
    if path in _PREP:
        return _PREP[path]
    try:
        im = Image.open(path)
        w, h = im.size
        if max(w, h) <= max_side:
            _PREP[path] = path
            return path
        sc = max_side / max(w, h)
        im = im.convert("RGBA") if im.mode in ("RGBA", "P", "LA") else im.convert("RGB")
        im = im.resize((max(1, int(w * sc)), max(1, int(h * sc))), Image.LANCZOS)
        ext = ".png" if im.mode == "RGBA" else ".jpg"
        outp = os.path.join(_TMP, hashlib.md5(os.path.abspath(path).encode("utf-8")).hexdigest() + ext)  # 固定檔名：重產結果可重現
        im.save(outp) if ext == ".png" else im.save(outp, quality=88)
        _PREP[path] = outp
        return outp
    except Exception:
        _PREP[path] = path
        return path


# ---------------- 基本形狀與文字 ----------------
def set_bg(slide, hexc):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = C(hexc)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    try:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.NONE
    except Exception:
        pass
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tb, tf


def _font(r, size, color, bold, italic=False):
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = C(color)
    if italic:
        f.italic = True
    rpr = r._r.get_or_add_rPr()
    rpr.set("lang", "zh-TW")
    rpr.set("altLang", "en-US")
    for tag in ("a:latin", "a:ea", "a:cs"):
        e = rpr.find(qn(tag))
        if e is None:
            e = rpr.makeelement(qn(tag), {})
            rpr.append(e)
        e.set("typeface", FONT)


def para(tf, text, size, color=None, bold=False, align=PP_ALIGN.LEFT, space_after=6, line=1.2,
         first=False, markup=True, prefix=None):
    """text 可含 [[藍:…]] 標記；prefix=(文字, 色) 例如條列圓點。"""
    color = color or G.D["ink"]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p._p.get_or_add_pPr().set("hangingPunct", "0")   # 全形標點不懸到欄外壓圖
    p.space_after = Pt(space_after)
    try:
        p.line_spacing = line
    except Exception:
        pass
    if prefix:
        r = p.add_run()
        r.text = prefix[0]
        _font(r, size, prefix[1], True)
    text = unicodedata.normalize("NFC", text or "")
    segs = parse_markup(text, G.EMPH) if markup else [(text, None)]
    for t, col in segs:
        if not t:
            continue
        r = p.add_run()
        r.text = t
        _font(r, size, col or color, bold)
    return p


def rect(slide, x, y, w, h, hexc=None, line_hex=None, line_w=1.0, rounded=False, radius=0.06, alpha=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if hexc:
        shp.fill.solid()
        shp.fill.fore_color.rgb = C(hexc)
        if alpha is not None:
            srgb = shp.fill.fore_color._xFill.find(qn('a:srgbClr'))
            srgb.append(srgb.makeelement(qn('a:alpha'), {'val': str(int(alpha * 1000))}))
    else:
        shp.fill.background()
    if line_hex:
        shp.line.color.rgb = C(line_hex)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def vbar(slide, x, y, h, hexc=None, w=0.085):
    return rect(slide, x, y, w, h, hexc=hexc or G.D["cobalt"])


# ---------------- 估算 ----------------
def em(text):
    return sum(1.0 if ord(ch) > 0x2E7F else 0.55 for ch in strip_markup(text))


_FONTS = {}
_TOK = re.compile(r"[A-Za-z0-9\-\.,;:%/()\[\]'’&+=<>µ·~_]+|\s+|.")
_WINFONTS = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
# 量字用的字型（名稱, Regular, Bold），依序取第一個存在的。PPTX 內的字型名稱一律是 common.FONT（微軟正黑體），
# 這裡只影響換行估算：Windows 用微軟正黑體本身；Ubuntu（雲端）用 fonts-noto-cjk 的 Noto Sans CJK TC（中文字寬相同，英文略寬）。
MEASURE_FONTS = [
    ("Microsoft JhengHei", os.path.join(_WINFONTS, "msjh.ttc"), os.path.join(_WINFONTS, "msjhbd.ttc")),
    ("Noto Sans CJK TC", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
     "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
    ("Noto Sans CJK TC", "/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc",
     "/usr/share/fonts/noto-cjk/NotoSansCJK-Bold.ttc"),
    ("Noto Sans CJK TC", "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Regular.ttc",
     "/usr/share/fonts/google-noto-cjk/NotoSansCJK-Bold.ttc"),
    ("Noto Sans CJK TC", "/usr/share/fonts/opentype/noto/NotoSansCJKtc-Regular.otf",
     "/usr/share/fonts/opentype/noto/NotoSansCJKtc-Bold.otf"),
]


def _tc_index(path):
    """.ttc 內繁中（TC）字面的 index；Noto Sans CJK 的 .ttc 依序含 JP／KR／SC／TC／HK。找不到就 0。"""
    if not path.lower().endswith(".ttc"):
        return 0
    for i in range(16):
        try:
            fam = ImageFont.truetype(path, 100, index=i).getname()[0] or ""
        except Exception:
            break
        if "TC" in fam.split() and "Mono" not in fam:
            return i
    return 0


def _pick_measure_font():
    """回傳 (名稱, Regular 路徑, Bold 路徑)；None＝用 em() 估算。
    環境變數 SLIDEKIT_MEASURE_FONT：em＝強制 em() 估算；字型檔路徑＝改用該字型（Bold 用 SLIDEKIT_MEASURE_FONT_BOLD，沒給就同一檔）。"""
    env = os.environ.get("SLIDEKIT_MEASURE_FONT", "").strip()
    if env.lower() == "em":
        return None
    if env:
        if os.path.isfile(env):
            bold = os.environ.get("SLIDEKIT_MEASURE_FONT_BOLD", "").strip()
            return (os.path.basename(env), env, bold if os.path.isfile(bold) else env)
        print(f"[warn] SLIDEKIT_MEASURE_FONT={env} 不存在，改用預設字型", file=sys.stderr)
    for name, reg, bold in MEASURE_FONTS:
        if os.path.isfile(reg):
            return (name, reg, bold if os.path.isfile(bold) else reg)
    return None


def measure_font_info():
    f = _pick_measure_font()
    return f"{f[0]}（{f[1]}）" if f else "無（em 估算：中文 1 字寬、英數 0.55；換行估算較粗）"


def _font_obj(bold):
    """量字寬的字型物件（微軟正黑體或 Noto Sans CJK TC，Regular／Bold）；找不到字型就回 None，改用 em() 估算。"""
    if bold not in _FONTS:
        _FONTS[bold] = None
        pick = _pick_measure_font()
        if pick:
            path = pick[2] if bold else pick[1]
            try:
                f = ImageFont.truetype(path, 100, index=_tc_index(path))
                try:  # 可變字型（例如 NotoSansTC-VF.ttf 預設是 Thin）切到 Regular／Bold；一般字型沒有 variation，略過
                    f.set_variation_by_name("Bold" if bold and pick[2] == pick[1] else "Regular")
                except Exception:
                    pass
                _FONTS[bold] = f
            except Exception:
                _FONTS[bold] = None
    return _FONTS[bold]


def _w(tok, bold):
    f = _font_obj(bold)
    return f.getlength(tok) / 100.0 if f is not None else em(tok)


def n_lines(text, size, width_in, bold=False):
    """以實際字型量測模擬換行：中文逐字可斷，英文單字（含相連括號、標點）不斷。"""
    cap = max(1.0, width_in * 72.0 / size)
    lines, cur = 1, 0.0
    for tok in _TOK.findall(strip_markup(text or "")):
        w = _w(tok, bold)
        if tok.isspace():
            if cur > 0:
                cur += w
            continue
        if cur + w <= cap:
            cur += w
        elif w > cap:
            total = cur + w
            lines += int(total // cap)
            cur = total % cap
        else:
            lines += 1
            cur = w
    return lines


def line_h(size, line=1.2):
    """估算行高（in）。1.12 係數依 PowerPoint 實際渲染 JhengHei 校正（1.05 每行少估約 0.07in、1.2 會誤報溢出；以 2026-09-25 免疫 deck 的 PNG 實測）。"""
    return size * line * 1.12 / 72.0


def note_fit(s, need, avail, what):
    G.layout.setdefault(s["id"], []).append(
        {"what": what, "need_in": round(need, 2), "avail_in": round(avail, 2), "overflow": need > avail + 0.05})


# ---------------- 圖片放置 ----------------
def place_image(slide, img, x, y, w, h, shadow=True):
    """圖片等比放進 (x,y,w,h)，下方留圖說；圖說以實際寬度估行數，反覆縮圖直到圖說不壓到左下「資料來源」。"""
    if not img or not img.get("exists"):
        return None
    try:
        iw, ih = Image.open(img["abspath"]).size
    except Exception:
        return None
    cred = img.get("src") or ""
    ctext = cred if cred.startswith(("圖", "資料來源")) else "圖：" + cred
    cw_of = lambda dw: max(dw, min(w, 4.5))
    ch = 0.0
    for _ in range(4):
        hh = h - ch
        floor = G.SH - getattr(G, "src_h", 0.0) - 0.10 - ch
        if y + hh > floor:
            hh = max(0.8, floor - y)
        far, iar = w / hh, iw / ih
        dw, dh = (w, w / iar) if iar > far else (hh * iar, hh)
        need = (n_lines(ctext, G.S["credit"], cw_of(dw)) * 0.25 + 0.06) if cred else 0.0
        if abs(need - ch) < 1e-6:
            break
        ch = need
    dx, dy = x + (w - dw) / 2, y
    pic = slide.shapes.add_picture(prep_image(img["abspath"]), Inches(dx), Inches(dy), Inches(dw), Inches(dh))
    if cred:
        cw = cw_of(dw)
        credit(slide, cred, x + (w - cw) / 2 if cw > dw else dx, cw, dy + dh + 0.04)
    return pic, dx, dy, dw, dh


def credit(slide, text, x, w, y):
    t = text if text.startswith(("圖", "資料來源")) else "圖：" + text
    lines = n_lines(t, G.S["credit"], w)
    _, tf = textbox(slide, x, y, w, 0.25 * lines)
    para(tf, t, G.S["credit"], color=G.D["grey"], align=PP_ALIGN.LEFT, first=True, line=1.0, markup=False)


def src_line(slide, srcs, dark=False):
    """文獻出處：左下 14pt「資料來源：…」，可兩行，不截斷。"""
    if not srcs:
        return
    text = "資料來源：" + "；".join(srcs)
    w = G.SW - 2 * G.M
    lines = n_lines(text, G.S["credit"], w)
    _, tf = textbox(slide, G.M, G.SH - 0.12 - 0.25 * lines, w, 0.25 * lines)
    para(tf, text, G.S["credit"], color=(DARK_TEXT["dim"] if dark else G.D["grey"]), first=True, line=1.0,
         markup=False)


def real_imgs(s):
    return [i for i in s.get("imgs", []) if i.get("exists")]


# ---------------- 標題 ----------------
def content_title(slide, s, size=None):
    t = s["title"]
    w = G.SW - 2 * G.M - 0.26
    size = size or G.S["title"]
    lines = n_lines(t, size, w, bold=True)
    if lines > 2 and size == G.S["title"]:
        size = G.S["title_small"]
        lines = n_lines(t, size, w, bold=True)
    h = lines * line_h(size, 1.1)
    vbar(slide, G.M, 0.36, h - 0.04)
    _, tf = textbox(slide, G.M + 0.26, 0.30, w, h + 0.1)
    para(tf, t, size, color=G.D["deep"], bold=True, first=True, line=1.1)
    return 0.30 + h + 0.30


# ---------------- 版型 ----------------
def L_title(slide, s):
    set_bg(slide, G.D["deep"])
    rect(slide, G.SW * 0.62, 0, G.SW * 0.38, G.SH, hexc=G.D["cobaltd"], alpha=22)
    vbar(slide, G.M + 0.1, 2.1, 1.9, hexc=G.D["terra"], w=0.10)
    _, tf = textbox(slide, G.M + 0.42, 2.0, G.SW - 2 * G.M - 0.6, 4.6)
    para(tf, s["title"], G.S["cover"], color="FFFFFF", bold=True, first=True, line=1.08, space_after=14)
    for i, b in enumerate(s["bullets"][:4]):
        para(tf, bullet_level(b)[1], G.S["cover_sub"] if i == 0 else G.S["min"],
             color=DARK_TEXT["mid"] if i == 0 else DARK_TEXT["low"], line=1.25, space_after=8)


def L_divider(slide, s, sec_label=""):
    set_bg(slide, G.D["deep"])
    rect(slide, G.SW * 0.64, 0, G.SW * 0.36, G.SH, hexc=G.D["cobaltd"], alpha=20)
    if sec_label and not s["title"].startswith(sec_label):
        _, ef = textbox(slide, G.M + 0.62, 1.95, G.SW - 2 * G.M, 0.5)
        para(ef, sec_label, G.S["divider_sub"], color=DARK_TEXT["low"], bold=True, first=True)
    vbar(slide, G.M + 0.32, 2.55, 1.8, hexc=G.D["terra"], w=0.10)
    _, tf = textbox(slide, G.M + 0.62, 2.5, G.SW - 2 * G.M - 0.7, 3.0)
    para(tf, s["title"], G.S["divider"], color="FFFFFF", bold=True, first=True, line=1.2)
    if s["bullets"]:
        _, bf = textbox(slide, G.M + 0.62, 5.3, G.SW - 2 * G.M - 0.7, 1.6)
        for i, b in enumerate(s["bullets"][:2]):
            para(bf, bullet_level(b)[1], G.S["divider_sub"], color=DARK_TEXT["mid"], first=(i == 0), line=1.2)


def bullets_need(bullets, size, w):
    need = 0.0
    for b in bullets:
        lvl, text = bullet_level(b)
        need += n_lines(text, size, w - (0.32 + 0.38 * lvl)) * line_h(size) + 8 / 72.0
    return need


def render_bullets(slide, s, bullets, x, y, w, h, size):
    _, tf = textbox(slide, x, y, w, h)
    need = 0.0
    for i, b in enumerate(bullets):
        lvl, text = bullet_level(b)
        ind = 0.32 + 0.38 * lvl
        mark = "• " if lvl == 0 else "– "
        p = para(tf, text, size, first=(i == 0), line=1.2, space_after=8,
                 prefix=(mark, G.D["cobalt"] if lvl == 0 else G.D["grey"]))
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(int(ind * 914400)))
        pPr.set('indent', str(int(-0.32 * 914400)))
        need += n_lines(text, size, w - ind) * line_h(size) + 8 / 72.0
    note_fit(s, need, h, "條列")


def L_left(slide, s):
    set_bg(slide, G.BG)
    y0 = content_title(slide, s)
    imgs = real_imgs(s)
    avail = G.SH - y0 - 0.45 - getattr(G, "src_h", 0.0)
    if not imgs:
        render_bullets(slide, s, s["bullets"], G.M + 0.1, y0, G.SW - 2 * G.M - 0.2, avail, G.S["body_noimg"])
        return
    tw = (G.SW - 2 * G.M) * 0.47
    if s["bullets"]:
        # 文字放不下就把文字欄加寬（47%→65%），圖讓位；仍放不下則由 note_fit 報出
        for frac in (0.47, 0.52, 0.57, 0.62, 0.65):
            tw = (G.SW - 2 * G.M) * frac
            if bullets_need(s["bullets"], G.S["body"], tw) <= avail:
                break
        render_bullets(slide, s, s["bullets"], G.M + 0.05, y0, tw, avail, G.S["body"])
        ix = G.M + tw + 0.40
    else:
        ix = G.M
    place_image(slide, imgs[0], ix, y0, G.SW - G.M - ix, avail - 0.15)


def _lead(slide, s, lines_max=None):
    """標題下的說明句（pair／flash／big 共用）；回傳圖區起點 y。"""
    y0 = content_title(slide, s)
    w = G.SW - 2 * G.M - 0.26
    need = 0.0
    if s["bullets"]:
        _, tf = textbox(slide, G.M + 0.26, y0, w, 2.4)
        for i, b in enumerate(s["bullets"] if lines_max is None else s["bullets"][:lines_max]):
            t = bullet_level(b)[1]
            para(tf, t, G.S["lead"], color=G.D["cobaltd"], first=(i == 0), line=1.2, space_after=4)
            need += n_lines(t, G.S["lead"], w) * line_h(G.S["lead"]) + 4 / 72.0
    return y0 + need + (0.12 if need else 0)


def L_flash(slide, s):
    set_bg(slide, G.BG)
    y = _lead(slide, s)
    imgs = real_imgs(s)
    if imgs:
        place_image(slide, imgs[0], G.M, y, G.SW - 2 * G.M, G.SH - y - 0.4)
    note_fit(s, y, G.SH - 2.0 if imgs else G.SH - 0.4, "說明句")


def L_big(slide, s):
    """單句大字：教師常用 36–50pt 一句話（「炎 = 火 + 火」）。"""
    set_bg(slide, G.BG)
    y0 = content_title(slide, s) if s["title"].strip() else 0.8
    imgs = real_imgs(s)
    w = G.SW - 2 * G.M - 0.26
    _, tf = textbox(slide, G.M + 0.26, y0 + 0.1, w, 3.0)
    need = 0.0
    for i, b in enumerate(s["bullets"][:3]):
        t = bullet_level(b)[1]
        para(tf, t, G.S["big"], color=G.D["deep"], bold=False, first=(i == 0), line=1.15, space_after=10)
        need += n_lines(t, G.S["big"], w) * line_h(G.S["big"], 1.15) + 10 / 72.0
    y = y0 + 0.1 + need + 0.15
    if imgs:
        place_image(slide, imgs[0], G.M, y, G.SW - 2 * G.M, G.SH - y - 0.4)
    note_fit(s, y, G.SH - (2.0 if imgs else 0.4), "大字")


def L_pair(slide, s):
    set_bg(slide, G.BG)
    y = _lead(slide, s)
    imgs = (real_imgs(s) + [None, None])[:2]
    gap = 0.3
    w = (G.SW - 2 * G.M - gap) / 2
    labs = [i.get("label") for i in imgs if i and i.get("label")]
    lab_h = (max(n_lines(t, G.S["label"], w, bold=True) for t in labs) * line_h(G.S["label"], 1.0) + 0.12) if labs else 0
    h = G.SH - y - 0.35 - lab_h
    for k, img in enumerate(imgs):
        x = G.M + k * (w + gap)
        r = place_image(slide, img, x, y + lab_h, w, h)
        if img and img.get("label"):
            _, tf = textbox(slide, x, y, w, lab_h)
            para(tf, img["label"], G.S["label"], color=G.D["cobaltd"], bold=True, align=PP_ALIGN.CENTER,
                 first=True, line=1.0)
    note_fit(s, y, G.SH - 2.2, "說明句")


def L_robbins(slide, s):
    set_bg(slide, G.BG)
    y = _lead(slide, s)
    imgs = real_imgs(s)
    if imgs:
        place_image(slide, imgs[0], G.M, y, G.SW - 2 * G.M, G.SH - y - 0.3)


def set_cell(cell, text, size, color=None, bold=False, fill=None, align=PP_ALIGN.LEFT):
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = C(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.10)
    cell.margin_right = Inches(0.08)
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    tf = cell.text_frame
    tf.word_wrap = True
    para(tf, text, size, color=color, bold=bold, align=align, first=True, line=1.1, space_after=0)


def L_table3(slide, s):
    set_bg(slide, G.BG)
    y0 = content_title(slide, s)
    tbl = s.get("table") or []
    if not tbl:
        return L_left(slide, s)
    if s["bullets"]:   # 全部說明句都畫（不再只畫第一條）
        lw = G.SW - 2 * G.M - 0.26
        need = sum(n_lines(bullet_level(b)[1], G.S["lead"], lw) * line_h(G.S["lead"]) + 4 / 72.0 for b in s["bullets"])
        _, tf = textbox(slide, G.M + 0.26, y0, lw, need + 0.1)
        for i, b in enumerate(s["bullets"]):
            para(tf, bullet_level(b)[1], G.S["lead"], color=G.D["cobaltd"], first=(i == 0), space_after=4)
        y0 += need + 0.1
    ncol = len(tbl[0])
    rows = [(r + [""] * ncol)[:ncol] for r in tbl]
    tw = G.SW - 2 * G.M
    w0 = max(em(r[0]) for r in rows)
    wr = [max(em(r[j]) for r in rows) for j in range(ncol)]
    tot = sum(max(4, x) for x in wr)
    widths = [tw * max(4, x) / tot for x in wr]
    widths = [max(1.3, x) for x in widths]
    scale = tw / sum(widths)
    widths = [x * scale for x in widths]
    avail = G.SH - y0 - 0.45
    size = G.S["table"]
    for cand in (G.S["table"], G.S["table_min"]):
        size = cand
        need = sum(max(n_lines(r[j], cand, widths[j] - 0.2, bold=(i == 0 or j == 0)) for j in range(ncol))
                   * line_h(cand, 1.1) + 0.14 for i, r in enumerate(rows))   # 0.14＝儲存格上下邊界＋框線
        if need <= avail:
            break
    note_fit(s, need, avail, "表格")
    t = slide.shapes.add_table(len(rows), ncol, Inches(G.M), Inches(y0), Inches(tw), Inches(min(avail, need))).table
    t.first_row = False
    t.horz_banding = False
    for j in range(ncol):
        t.columns[j].width = Inches(widths[j])
    for i, r in enumerate(rows):
        for j, c in enumerate(r):
            if i == 0:
                set_cell(t.cell(i, j), c, size, color="FFFFFF", bold=True, fill=G.D["cobaltd"])
            else:
                set_cell(t.cell(i, j), c, size, color=(G.D["terrad"] if j == 0 else G.D["ink"]), bold=(j == 0),
                         fill=("FFFFFF" if i % 2 else G.D["tint"]))


def quiz_title(course, deck, q):
    tpl = course.deck(deck).get("quiz_label", "{年度}-{第幾次} 第{題號}題")
    return tpl.format(年度=q["年度"], 第幾次=CN_NUM.get(int(q["第幾次"]), q["第幾次"]), 題號=q["題號"])


def quiz_style(course, deck):
    """course.yaml decks.<deck>.quiz_style：乙（預設；年次當標題、正解選項整段強調色）或 甲（1 列 3 欄表）。"""
    st = str(course.deck(deck).get("quiz_style") or "乙").strip()
    if st not in ("甲", "乙"):
        raise ValueError(f"decks.{deck}.quiz_style 只能是 甲 或 乙，現在是 {st!r}")
    return st


def L_quiz_jia(slide, s, q):
    """甲式（34_國考題 EXAM-03／04、11 LAY-20）：大堂課不放標題；1 列 3 欄表
    「正解字母（粗體、不上色）｜題號．題幹＋選項｜年度-次」，表寬＝畫幅扣左右邊界。"""
    W = G.SW - 2 * G.M
    c1, c3 = 0.8, 1.3
    c2 = W - c1 - c3
    y0 = 0.6
    avail = G.SH - y0 - 0.45 - G.src_h
    ans = ",".join(a.strip() for a in str(q["answer"]).split(","))
    opts = [(k, q["options"][k]) for k in "ABCDE" if k in q["options"]]
    stem = f"{q['題號']}. {q['question']}"
    tw = c2 - 0.2
    size, need = G.S["quiz_min"], 0.0
    for cand in range(G.S["quiz"], G.S["quiz_min"] - 1, -2):
        need = n_lines(stem, cand, tw) * line_h(cand, 1.1)
        need += sum(n_lines(f"({k}) {v}", cand, tw) * line_h(cand, 1.1) for k, v in opts) + 0.2
        size = cand
        if need <= avail:
            break
    note_fit(s, need, avail, "國考題（甲式）")
    shape = slide.shapes.add_table(1, 3, Inches(G.M), Inches(y0), Inches(W), Inches(min(avail, need)))
    t = shape.table
    t.first_row = False
    t.horz_banding = False
    for j, w in enumerate((c1, c2, c3)):
        t.columns[j].width = Inches(w)
    set_cell(t.cell(0, 0), ans, size, color=G.D["ink"], bold=True, fill="FFFFFF", align=PP_ALIGN.CENTER)
    set_cell(t.cell(0, 1), stem, size, color=G.D["ink"], fill="FFFFFF")
    tf = t.cell(0, 1).text_frame
    for k, v in opts:
        para(tf, f"({k}) {v}", size, color=G.D["ink"], line=1.1, space_after=0, markup=False)
    set_cell(t.cell(0, 2), f"{q['年度']}-{q['第幾次']}", size, color=G.D["ink"], fill="FFFFFF", align=PP_ALIGN.CENTER)
    if "#" in str(q["answer"]):
        _, tf2 = textbox(slide, G.M, y0 + min(avail, need) + 0.1, W, 0.5)
        para(tf2, "（本題一律給分）", G.S["label"], color=G.D["grey"], first=True, line=1.0)


def L_quiz(slide, s, qs):
    set_bg(slide, G.BG)
    q = qs.get(s["quiz"] or "")
    if not q:
        y0 = content_title(slide, s)
        _, tf = textbox(slide, G.M, y0, 9, 1)
        para(tf, f"（題庫查無 {s['quiz']}）", G.S["body"], color=G.D["terrad"], first=True)
        return
    if quiz_style(G.course, s["deck"]) == "甲":
        return L_quiz_jia(slide, s, q)
    s = dict(s)
    s["title"] = s["title"] if s["title"].strip() and s["title"] != "國考題" else quiz_title(G.course, s["deck"], q)
    y0 = content_title(slide, s, size=G.S["title_small"])
    w = G.SW - 2 * G.M - 0.1
    avail = G.SH - y0 - 0.35
    ans = [a.strip() for a in str(q["answer"]).split(",")]
    opts = [(k, q["options"][k]) for k in "ABCDE" if k in q["options"]]
    size = G.S["quiz_min"]
    for cand in range(G.S["quiz"], G.S["quiz_min"] - 1, -2):
        need = n_lines(q["question"], cand, w) * line_h(cand, 1.15) + 0.15
        need += sum(n_lines(f"({k}) {v}", cand, w - 0.3) * line_h(cand, 1.15) + 0.06 for k, v in opts)
        size = cand
        if need <= avail:
            break
    note_fit(s, need, avail, "國考題")
    _, tf = textbox(slide, G.M + 0.05, y0, w, avail)
    para(tf, f"{q['題號']}. {q['question']}", size, color=G.D["deep"], first=True, line=1.15, space_after=10,
         markup=False)
    for k, v in opts:
        is_ans = k in ans
        p = para(tf, f"({k}) {v}", size, color=(G.EMPH.get("藍", "143F9C") if is_ans else G.D["ink"]),
                 line=1.15, space_after=4, markup=False)
        pPr = p._p.get_or_add_pPr()
        pPr.set('marL', str(int(0.45 * 914400)))
        pPr.set('indent', str(int(-0.45 * 914400)))
    if "#" in ans:
        para(tf, "（本題一律給分）", G.S["label"], color=G.D["grey"], line=1.0)
    if s.get("quiz_note"):
        para(tf, s["quiz_note"], G.S["label"], color=G.D["cobaltd"], line=1.1, space_after=0)


def L_stats(slide, s):
    set_bg(slide, G.BG)
    y0 = content_title(slide, s)
    tiles = [b.split("｜", 1) for b in s["bullets"] if "｜" in b]
    plain = [b for b in s["bullets"] if "｜" not in b]
    if tiles:
        n = len(tiles)
        gap = 0.3
        cw = (G.SW - 2 * G.M - (n - 1) * gap) / n
        ch = 2.4
        cols = [G.D["cobalt"], G.D["terra"], G.D["cobaltd"], G.D["terrad"]]
        for i, (big, label) in enumerate(tiles):
            x = G.M + i * (cw + gap)
            rect(slide, x, y0 + 0.1, cw, ch, hexc=G.D["tint"], line_hex=G.D["line"], rounded=True)
            _, tf = textbox(slide, x + 0.12, y0 + 0.1, cw - 0.24, ch, anchor=MSO_ANCHOR.MIDDLE)
            para(tf, big.strip(), G.S["stats_big"], color=cols[i % 4], bold=True, align=PP_ALIGN.CENTER,
                 first=True, line=1.0, space_after=6)
            para(tf, label.strip(), G.S["label"], align=PP_ALIGN.CENTER, line=1.1)
        y0 += ch + 0.3
    if plain:
        render_bullets(slide, s, plain, G.M + 0.1, y0, G.SW - 2 * G.M - 0.2, G.SH - y0 - 0.45, G.S["body_noimg"])


def _cell_text(tf, text, size, color, bold=False, align=PP_ALIGN.CENTER):
    """格子文字：以換行分段（中文一行、英文一行），同字級。"""
    for k, part in enumerate(str(text).split(NL)):
        para(tf, part, size, color=color, bold=bold, align=align, first=(k == 0), line=1.0, space_after=0)


def _cell_need(text, size, w):
    return sum(n_lines(part, size, w) for part in str(text).split(NL)) * line_h(size, 1.0) + 0.18


def L_diagram(slide, s):
    """框架圖：course.yaml diagrams.<名>：nodes [{label, group}]、rows [{name, cells[]}]、groups {g:{main,light}}。
    orient: cols（預設；節點橫排、屬性為列）或 rows（節點直排、屬性為欄，適合 5 個節點＋長雙語字）。
    格子文字可用換行字元分行（中文一行、英文一行；course.yaml 裡寫成 YAML 的 "\\n"）。點亮語彙：全部／節點 label／group 名／列名；未點亮畫淺灰。"""
    set_bg(slide, G.BG)
    y0 = content_title(slide, s)
    spec = G.course["diagrams"].get(s["diagram"])
    if not spec:
        _, tf = textbox(slide, G.M, y0, 9, 1)
        para(tf, f"（course.yaml 沒有框架圖 {s['diagram']}）", G.S["body"], color=G.D["terrad"], first=True)
        return
    bright = set(s["diagram_bright"] or ["全部"])
    allon = "全部" in bright
    nodes, rows, groups = spec["nodes"], spec.get("rows", []), spec.get("groups", {})
    lit = lambda nd: allon or nd["label"].replace(NL, "") in bright or nd["label"] in bright or nd.get("group") in bright
    gcol = lambda nd: groups.get(nd.get("group"), {"main": G.D["cobalt"], "light": "E3EDF8"})
    cs = G.S["min"]
    lab = G.S["label"]
    gap = 0.10
    avail = G.SH - y0 - 0.4
    W = G.SW - 2 * G.M

    def box(x, y, w, h, text, size, on, g, bold=False, head=False):
        fill = g["light"] if on else "FFFFFF"
        line = g["main"] if on else ("C3D3E0" if head else "E2EAF1")
        rect(slide, x, y, w, h, hexc=fill, line_hex=line, line_w=(1.8 if head else 1.0), rounded=True,
             radius=(0.14 if head else 0.10))
        _, tf = textbox(slide, x + 0.07, y + 0.03, w - 0.14, h - 0.06, anchor=MSO_ANCHOR.MIDDLE)
        col = (G.D["deep"] if head else G.D["ink"]) if on else ("8AA0B5" if head else "BCCBD8")
        _cell_text(tf, text, size, col, bold=(bold and on))
        need = _cell_need(text, size, w - 0.14)
        if need > h + 0.05:
            note_fit(s, need, h, f"框架圖格子「{str(text)[:10]}」")

    def fit_heights(needs, space, cap):
        """先給每列需要的高度，剩下的平均分（每列不超過 cap）；放不下就等比縮並記錄溢出。"""
        tot = sum(needs)
        if tot > space:
            note_fit(s, tot, space, "框架圖總高")
            return [h * space / tot for h in needs]
        extra = (space - tot) / len(needs)
        return [min(max(h, cap), h + extra) if h < cap else h for h in needs]

    if spec.get("orient") == "rows":
        nc = len(rows)
        w0 = spec.get("node_w", 3.0)
        cw = spec.get("col_w") or [(W - w0 - nc * gap) / nc] * nc
        k = (W - w0 - nc * gap) / sum(cw)
        cw = [x * k for x in cw]
        xs = [G.M + w0 + gap + sum(cw[:j]) + j * gap for j in range(nc)]
        hh = 0.42
        for j, r in enumerate(rows):
            _, tf = textbox(slide, xs[j], y0, cw[j], hh, anchor=MSO_ANCHOR.MIDDLE)
            para(tf, r["name"], cs, color=G.D["deep"], bold=True, align=PP_ALIGN.CENTER, first=True, line=1.0,
                 space_after=0)
        y = y0 + hh + 0.05
        space = avail - hh - 0.05 - gap * (len(nodes) - 1)
        for labsz in (lab, cs):
            need = [max([_cell_need(nd["label"], labsz, w0 - 0.14)] +
                        [_cell_need(r["cells"][i], cs, cw[j] - 0.14) for j, r in enumerate(rows)])
                    for i, nd in enumerate(nodes)]
            if sum(need) <= space:
                break
        hs = fit_heights(need, space, 1.2)
        for i, nd in enumerate(nodes):
            on = lit(nd)
            box(G.M, y, w0, hs[i], nd["label"], labsz, on, gcol(nd), bold=True, head=True)
            for j, r in enumerate(rows):
                on_c = on and (allon or r["name"] in bright)
                box(xs[j], y, cw[j], hs[i], r["cells"][i], cs, on_c, gcol(nd))
            y += hs[i] + gap
        return

    lbl_w = spec.get("label_w", 0.95) if rows else 0.0
    n = len(nodes)
    nw = (W - lbl_w - (n - 1) * gap) / n
    x0 = G.M + lbl_w
    nh = max(0.8, max(_cell_need(nd["label"], lab, nw - 0.14) for nd in nodes))
    for i, nd in enumerate(nodes):
        box(x0 + i * (nw + gap), y0, nw, nh, nd["label"], lab, lit(nd), gcol(nd), bold=True, head=True)
    y = y0 + nh + 0.12
    left = avail - nh - 0.12
    need = [max(_cell_need(c, cs, nw - 0.14) for c in r["cells"]) for r in rows]
    hs = fit_heights(need, left - gap * (len(rows) - 1), 1.6)
    for r, rh in zip(rows, hs):
        on_row = allon or r["name"] in bright
        _, tf = textbox(slide, G.M, y, lbl_w - 0.08, rh, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, r["name"], cs, color=(G.D["deep"] if on_row else "A8BACA"), bold=True,
             align=PP_ALIGN.RIGHT, first=True, line=1.0, space_after=0)
        for i, cell in enumerate(r["cells"]):
            nd = nodes[i]
            box(x0 + i * (nw + gap), y, nw, rh, cell, cs, on_row and lit(nd), gcol(nd))
        y += rh + gap


RENDER = {"title": L_title, "left": L_left, "flash": L_flash, "big": L_big, "pair": L_pair,
          "robbins": L_robbins, "table3": L_table3, "stats": L_stats, "diagram": L_diagram}


def build_pptx(course, deck, out_path, section=None):
    setup(course)
    G.layout = {}
    pool = parse_pool(course)
    built = build_deck(course, deck, pool=pool)
    if section:
        lo, hi, t = built["sections"][section - 1]
        built["slides"] = [s for s in built["slides"] if lo <= s["num"] <= hi]
        for k, s in enumerate(built["slides"], 1):
            s["num"] = k
        built["sections"] = [(1, len(built["slides"]), t)]
    notes = notes_map(built, parse_script(course))
    prs = Presentation()
    prs.slide_width = Inches(G.SW)
    prs.slide_height = Inches(G.SH)
    blank = prs.slide_layouts[6]
    errors = []
    for s in built["slides"]:
        slide = prs.slides.add_slide(blank)
        srcs = s.get("src") or []
        G.src_h = (n_lines("資料來源：" + "；".join(srcs), G.S["credit"], G.SW - 2 * G.M) * 0.25 + 0.05) if srcs else 0.0
        try:
            if s["type"] == "divider":
                L_divider(slide, s, sec_label=(s["section"] if s.get("section_first") else ""))
            elif s["type"] == "quiz":
                L_quiz(slide, s, built["questions"])
            else:
                RENDER.get(s["type"], L_left)(slide, s)
        except Exception as e:  # 單張錯誤不讓整套失敗，check_deck 會列出
            set_bg(slide, "FFFFFF")
            _, tf = textbox(slide, 0.6, 0.6, 8.8, 5)
            para(tf, f"[{s['type']}] #{s['num']} {s['id']} render error: {e}", 20, color="B00020", first=True,
                 markup=False)
            errors.append((s["num"], s["id"], repr(e)))
        if s["type"] not in ("title",):
            src_line(slide, s.get("src"), dark=(s["type"] == "divider"))
        nt = notes.get(s["id"], "")
        if nt:
            slide.notes_slide.notes_text_frame.text = nt
    prs.save(out_path)
    rep = {s["id"]: {"num": s["num"], "fits": G.layout.get(s["id"], [])} for s in built["slides"]}
    json.dump({"errors": errors, "measure_font": measure_font_info(), "slides": rep}, open(os.path.splitext(out_path)[0] + "_layout.json", "w",
                                                       encoding="utf-8", newline="\n"), ensure_ascii=False, indent=1)
    return built, errors


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--course", required=True)
    ap.add_argument("--deck", default="all")
    ap.add_argument("--date", default=None)
    ap.add_argument("--section", type=int, default=None)
    ap.add_argument("--name", default=None)
    ap.add_argument("--no-check", action="store_true")
    a = ap.parse_args()
    course = load_course(a.course)
    decks = course["deck_order"] if a.deck == "all" else [a.deck]
    out_dir = course.path("out_dir")
    os.makedirs(out_dir, exist_ok=True)
    ok = True
    print("量字字型：" + measure_font_info())
    import check_deck
    for d in decks:
        name = a.name or deck_filename(course, d, a.date)
        if a.name and len(decks) > 1:
            name += f"_{d}"
        if a.section:
            name += f"_第{a.section}節"
        outp = os.path.join(out_dir, name + ".pptx")
        if os.path.exists(outp):
            print(f"[注意] 覆寫既有檔：{outp}（若是老師改過的版本，先照 SOP 第 11b 步回灌再重產）")
        built, errors = build_pptx(course, d, outp, section=a.section)
        print(f"OK {built['name']} {len(built['slides'])} 張 -> {outp}")
        for num, sid, e in errors:
            print(f"  render error #{num} {sid}: {e}")
            ok = False
        if not a.no_check:
            ok = check_deck.main(course, outp, d, section=a.section) and ok
    return 0 if ok else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
